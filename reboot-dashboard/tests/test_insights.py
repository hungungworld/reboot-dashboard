"""시장 · 고객 인사이트 — 규칙 엔진 · 1회분 단가 · 수집기(가짜 응답) · 데이터 업데이트 · 화면 API · 내려받기 (인터넷 · 실제 키 없이)."""
import io
import json
from datetime import date, timedelta

import pytest

from server.insights import http, rules, service, update
from server.insights.util import bucket, parse_servings, review_month, week_key

TODAY = date(2026, 10, 23)


class FakeRes:
    def __init__(self, status=200, data=None):
        self.status_code = status
        self._data = data

    def json(self):
        if self._data is None:
            raise ValueError("no json")
        return self._data


# ───────── 가짜 외부 응답 ─────────

def _weekly_points(spike=True):
    """2025-10-26부터 52주, 평균 40 근처 · 작년 핼러윈 다음 주(2025-11-02)에 90."""
    start = date(2025, 10, 26)
    out = []
    for i in range(52):
        d = start + timedelta(weeks=i)
        v = 40 + (i % 5)
        if spike and d == date(2025, 11, 2):
            v = 90
        out.append({"date": d.strftime("%b %d, %Y"), "timestamp": str(int((d - date(1970, 1, 1)).total_seconds())),
                    "values": [{"query": "hangover", "value": str(v), "extracted_value": v}]})
    return {"interest_over_time": {"timeline_data": out}}


SEARCH = {
    "hangover": {"organic_results": [
        {"asin": "B0SPONSOR1", "title": "Sponsored hangover thing", "sponsored": True, "extracted_price": 9.99},
        {"asin": "B0COMP0001", "title": "Flyby Hangover Recovery Aid, 30 Count", "price": "$29.99", "extracted_price": 29.99,
         "rating": 4.3, "reviews": 5400, "link": "https://www.amazon.com/dp/B0COMP0001"},
        {"asin": "B0COMP0002", "title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0, "rating": 4.1,
         "reviews": 120},
        {"asin": "B0COMP0003", "title": "Haejang soup instant kit", "extracted_price": 12.0, "reviews": 30},
    ]},
    "korean hangover": {"organic_results": [
        {"asin": "B0COMP0002", "title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0, "reviews": 120},
        {"asin": "B0COMP0004", "title": "Liquid I.V. Hydration Multiplier - 16 Count", "extracted_price": 25.0, "reviews": 90000},
    ]},
}
PRODUCTS = {
    "B0COMP0001": {"product_results": {"title": "Flyby Hangover Recovery Aid, 30 Count", "brand": "Visit the Flyby Store",
                                       "extracted_price": 29.99, "stock": "In Stock", "rating": 4.3, "reviews": 5400},
                   "about_item": ["Supports recovery", "30 capsules"],
                   "reviews_information": {"summary": {"text": "Customers like the convenience."},
                                           "authors_reviews": [{"title": "Works", "text": "It worked great after a party, easy to carry.",
                                                                "rating": 5, "date": "October 3, 2026"}]}},
    "B0COMP0004": {"product_results": {"title": "Liquid I.V. Hydration Multiplier - 16 Count", "brand": "Liquid I.V.",
                                       "extracted_price": 25.0, "stock": "In Stock"}},
    "B0COMP0002": {"product_results": {"title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0}},
    "B0OWNASIN1": {"product_results": {"title": "Aldicom Hangover Jelly 1 pack*5 pack", "extracted_price": 15.0,
                                       "stock": "Currently unavailable."},
                   "reviews_information": {"authors_reviews": [
                       {"title": "Bad taste", "text": "The taste is too bitter and the flavor is gross.", "rating": 2},
                       {"title": "Meh", "text": "Bitter aftertaste, did not like the taste.", "rating": 1}]}},
}
WEEE_MD = """# Hangover

[![Condition Hangover Drink](https://img.weeecdn.com/a.jpg)](https://www.sayweee.com/en/product/Condition-Hangover-Drink/12345)

[Condition Hangover Drink 100ml x 10](https://www.sayweee.com/en/product/Condition-Hangover-Drink/12345)

$19.99$24.99

[![Aldicom Jelly](https://img.weeecdn.com/b.jpg)](https://www.sayweee.com/en/product/Aldicom-Jelly/67890)

[Aldicom Hangover Jelly 1 pack*5 pack](https://www.sayweee.com/en/product/Aldicom-Jelly/67890)

$12.99

Sold out

[Haejang Soup](https://www.sayweee.com/en/product/Haejang-Soup/11111)

$8.99
"""
YT_SEARCH = {"items": [{"id": {"videoId": "vid00000001"}, "snippet": {"title": "Korean hangover drink review", "channelTitle": "K-Snack Lab",
                                                                    "channelId": "UC123", "publishedAt": "2026-09-01T00:00:00Z"}},
                       {"id": {"videoId": "vid00000002"}, "snippet": {"title": "Small video", "channelTitle": "Tiny",
                                                                    "channelId": "UC999", "publishedAt": "2026-09-02T00:00:00Z"}}]}
YT_STATS = {"items": [{"id": "vid00000001", "statistics": {"viewCount": "50000", "likeCount": "900", "commentCount": "4"}},
                      {"id": "vid00000002", "statistics": {"viewCount": "500", "commentCount": "1"}}]}
YT_COMMENTS = {"items": [
    {"snippet": {"topLevelComment": {"id": f"c{i}", "snippet": {"textOriginal": t, "likeCount": 5 - i,
                                                                 "publishedAt": "2026-09-03T00:00:00Z"}}}}
    for i, t in enumerate(["Where can I buy this in the US?", "Where to buy the korean hangover drink online?",
                           "Does it actually work for a bad hangover?", "great video bro"])]}


@pytest.fixture()
def fake_net(monkeypatch):
    calls = []

    def fake_get(url, params=None, headers=None, timeout=25):
        p = params or {}
        calls.append(("GET", url, dict(p)))
        if "serpapi.com" in url:
            if p.get("api_key") == "bad-serp-key-000":
                return FakeRes(401, {"error": "Invalid API key. Your API key should be here: https://serpapi.com/manage-api-key"})
            eng = p.get("engine")
            if eng == "amazon":
                return FakeRes(200, SEARCH.get(p["k"], {"organic_results": []}))
            if eng == "amazon_product":
                return FakeRes(200, PRODUCTS.get(p["asin"], {"error": "Google hasn't returned any results for this query."}))
            if eng == "google_trends":
                return FakeRes(200, _weekly_points(spike=p["q"] == "hangover"))
        if "googleapis.com/youtube" in url:
            if url.endswith("/search"):
                return FakeRes(200, YT_SEARCH)
            if url.endswith("/videos"):
                return FakeRes(200, YT_STATS)
            if url.endswith("/commentThreads"):
                if p.get("videoId") == "vid00000001":
                    return FakeRes(200, YT_COMMENTS)
                return FakeRes(403, {"error": {"errors": [{"reason": "commentsDisabled"}]}})
        raise AssertionError(f"예상하지 못한 호출: {url}")

    def fake_post(url, json=None, headers=None, timeout=90):
        calls.append(("POST", url, dict(json or {})))
        if "firecrawl.dev" in url:
            assert headers["Authorization"].startswith("Bearer ")
            return FakeRes(200, {"success": True, "data": {"markdown": WEEE_MD}})
        raise AssertionError(f"예상하지 못한 호출: {url}")

    monkeypatch.setattr(http, "get", fake_get)
    monkeypatch.setattr(http, "post", fake_post)
    return calls


@pytest.fixture()
def fixed_today(monkeypatch):
    monkeypatch.setattr(service, "today", lambda: TODAY)
    return TODAY


@pytest.fixture()
def sync_jobs(monkeypatch):
    """데이터 업데이트를 뒤가 아니라 바로 실행 (테스트용).
    jobs.run_background는 시작하면 Thread, 이미 실행 중이면 None을 돌려주므로 None이 아닌 값을 돌려줌."""
    def run_now(app, name, fn, *a, **k):
        fn(*a, **k)
        return "started"
    monkeypatch.setattr(update.jobs, "run_background", run_now)


def _all_keys(monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "serp-test-key-1234")
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-test-key-1234567890")
    monkeypatch.setenv("YOUTUBE_API_KEY", "AIza-test-key-5678")


def _own_asin_config(client):
    cfg = client.get("/api/settings/collect").get_json()["config"]
    cfg["tracked_asins"] = [{"asin": "B0OWNASIN1", "label": "알디콤 젤리", "own": True}]
    cfg["trend_keywords"] = ["hangover", "hangover drink"]
    cfg["search_sets"] = {"A": ["korean hangover drink review"], "B": [], "K": []}
    assert client.put("/api/settings/collect", json={"config": cfg}).status_code == 200


# ───────── 1회분 단가 ─────────

@pytest.mark.parametrize("text, n", [
    ("5 pack*4 boxes", 20), ("Pack of 6", 6), ("Hangover jelly 10 sticks (min 2)", 10), ("1 pack*5 pack", 5),
    ("21 ml × 4 Sachets", 4), ("숙취해소 젤리 20팩", 20), ("12 Count (Pack of 2)", 24), ("Morning Recovery 3.4 fl oz, 6 Pack", 6),
    ("(min 2) hangover drink", None), ("DHM 1000 mg", None),
])
def test_parse_servings(text, n):
    assert parse_servings(text)[0] == n


def test_buckets_and_month_rules():
    assert [bucket(2.99), bucket(3.0), bucket(5.99), bucket(6.0), bucket(None)] == ["low", "mid", "mid", "high", None]
    assert review_month(date(2026, 11, 5)) == "2026-10" and review_month(date(2026, 11, 8)) == "2026-11"
    assert week_key(TODAY) == "2026-W43"


# ───────── 규칙 엔진 (화면과 같은 모양의 고정 데이터) ─────────

def _item(key, title, price, source="amazon", own=False, in_stock=True):
    return {"key": key, "title": title, "price": price, "source": source, "own": own, "in_stock": in_stock}


def _ctx(**kw):
    ctx = {"today": TODAY, "shelf_now": [], "shelf_prev": None, "trends": {}, "questions": [], "own_complaints": [],
           "seasons": [{"name": "Halloween", "start": "2026-10-31", "monitor": True},
                       {"name": "Thanksgiving Eve", "start": "2026-11-25", "monitor": True},
                       {"name": "Super Bowl", "start": "2027-02-14", "monitor": True}]}
    ctx.update(kw)
    return ctx


def test_rules_no_data_means_no_actions():
    assert rules.weekly_actions(_ctx()) == []


def test_rules_own_out_of_stock_needs_season_within_45_days():
    now = [_item("amazon:OWN", "Aldicom Jelly", 15, own=True, in_stock=False)]
    acts = rules.weekly_actions(_ctx(shelf_now=now))
    assert [a["rule"] for a in acts] == ["own_oos_season"] and acts[0]["grade"] == "긴급"
    assert "Halloween" in acts[0]["why"] and acts[0]["facts"]["남은 날"] == 8
    far = _ctx(shelf_now=now, seasons=[{"name": "Super Bowl", "start": "2027-02-14", "monitor": True}])
    assert rules.weekly_actions(far) == []
    off = _ctx(shelf_now=now, seasons=[{"name": "Halloween", "start": "2026-10-31", "monitor": False}])
    assert rules.weekly_actions(off) == []


def test_rules_price_changes_and_new_entries():
    prev = [_item("amazon:A", "Flyby 30 Count", 30.0), _item("amazon:B", "Liquid IV 16 Count", 25.0),
            _item("amazon:C", "Jelly 10 sticks", 20.0), _item("amazon:OWN", "Aldicom", 15.0, own=True)]
    now = [_item("amazon:A", "Flyby 30 Count", 24.0),            # -20% → 인하 긴급
           _item("amazon:B", "Liquid IV 16 Count", 28.0),         # +12% → 인상 참고
           _item("amazon:C", "Jelly 10 sticks", 21.0),            # +5% → 없음
           _item("amazon:OWN", "Aldicom", 9.0, own=True),         # 자사 가격 변동은 경쟁사 규칙에서 제외
           _item("amazon:D", "New Hangover Patch", 12.0),         # 신규
           _item("weee:1", "Weee item", 10.0, source="weee")]     # Weee!는 지난 묶음이 없어 비교 안 함
    acts = {a["rule"]: a for a in rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=prev))}
    assert set(acts) == {"price_cut", "price_up", "new_entry"}
    assert acts["price_cut"]["grade"] == "긴급" and "-20%" in acts["price_cut"]["fact"]
    assert acts["price_up"]["grade"] == "참고" and "+12%" in acts["price_up"]["fact"]
    assert acts["new_entry"]["grade"] == "참고" and "1개" in acts["new_entry"]["fact"]
    order = [a["grade"] for a in rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=prev))]
    assert order == sorted(order, key=lambda g: rules.GRADE_ORDER[g])
    # 첫 주(지난 묶음 없음)는 비교하지 않음
    assert rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=None)) == []


def test_rules_season_search_lift_questions_and_complaints():
    pts = [{"date": (date(2025, 10, 26) + timedelta(weeks=i)).isoformat(), "value": 90 if i == 1 else 40} for i in range(52)]
    lift = rules.season_lift(pts, "2026-10-31")
    assert lift and lift["date"] == "2025-11-02" and lift["pct"] > 100
    flat = [dict(p, value=40) for p in pts]
    assert rules.season_lift(flat, "2026-10-31")["pct"] == 0
    ctx = _ctx(trends={"hangover": pts, "flat": flat},
               questions=[{"label": "어디서 사나요 (구매처)", "count": 7, "examples": ["Where can I buy this?"]}],
               own_complaints=[{"topic": "맛·향", "count": 3, "examples": ["too bitter"]}])
    acts = rules.weekly_actions(ctx)
    by = {a["rule"]: a for a in acts}
    assert by["season_search"]["grade"] == "기회" and "'hangover'" in by["season_search"]["fact"]
    assert by["top_question"]["grade"] == "기회" and "7건" in by["top_question"]["fact"]
    assert by["own_complaint"]["grade"] == "긴급" and "맛·향" in by["own_complaint"]["fact"]
    assert acts[0]["rule"] == "own_complaint"  # 긴급이 먼저
    # Thanksgiving Eve(33일 뒤)는 30일 밖이라 기회 아님
    assert all("Thanksgiving" not in a["fact"] for a in acts)


# ───────── 화면 API (빈 상태) ─────────

def test_insights_requires_login(client):
    assert client.get("/api/insights/weekly").status_code == 401


def test_insights_page_loads_part_files(auth_client):
    html = auth_client.get("/dashboard/insights").get_data(as_text=True)
    assert "/parts/insights/insights.js" in html and "/parts/insights/insights.css" in html


def test_empty_states(auth_client, fixed_today):
    w = auth_client.get("/api/insights/weekly").get_json()
    assert w["hasData"] is False and w["actions"] == [] and w["week"] == "2026-W43"
    assert w["seasons"][0]["name"] == "Halloween" and w["seasons"][0]["dday"] == 8
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["month"] == "2026-10" and m["summary"] == [] and m["strategy"] is None
    assert m["metrics"]["ownOosWeeks"] is None and m["metrics"]["priceChanges"] is None
    for tab in ("shelf", "competitors", "reviews", "audience", "trends"):
        assert auth_client.get(f"/api/insights/data/{tab}").status_code == 200, tab
    assert auth_client.get("/api/insights/data/nope").status_code == 404
    assert auth_client.get("/api/insights/trends").get_json() == {"keywords": [], "updatedAt": None}
    assert auth_client.get("/api/insights/personas").get_json() == []
    assert auth_client.get("/api/insights/monthly?month=2026-13").status_code == 400


def test_update_plan_reports_missing_keys(auth_client):
    plan = auth_client.get("/api/insights/update/plan").get_json()
    by = {i["id"]: i for i in plan["items"]}
    assert by["amazon"]["ready"] is False and by["amazon"]["missing"] == ["SERPAPI_KEY"]
    assert by["reviews"]["ready"] is True and "규칙" in by["reviews"]["note"]
    assert by["monthly"]["ready"] is False
    assert by["amazon"]["cost"]["serpapi"] == 2 + 6  # 검색어 2 + 상세 6
    assert plan["limits"]["serpapi"] == 250
    res = auth_client.post("/api/insights/update", json={"items": ["amazon"]})
    assert res.status_code == 400 and "SerpApi" not in res.get_json()["error"] and "Amazon" in res.get_json()["error"]
    assert auth_client.post("/api/insights/update", json={"items": []}).status_code == 400
    assert auth_client.post("/api/insights/update", json={"items": ["hack"]}).status_code == 400


# ───────── 데이터 업데이트 (가짜 응답) ─────────

def test_update_job_collects_everything(auth_client, monkeypatch, fake_net, fixed_today, sync_jobs):
    _all_keys(monkeypatch)
    _own_asin_config(auth_client)
    items = ["weee", "amazon", "trends", "youtube", "reviews", "voc", "actions"]
    res = auth_client.post("/api/insights/update", json={"items": items})
    assert res.status_code == 202, res.get_json()
    st = auth_client.get("/api/insights/update/status").get_json()
    run = st["run"]
    assert run["status"] == "done", run["results"]
    assert set(run["results"]) == set(items)
    assert run["usage"]["serpapi"] >= 2 + 2 and run["usage"]["firecrawl"] == 1 and run["usage"]["youtube"] == 100 + 1 + 1
    # 키가 결과 · 로그에 남지 않음
    logs = auth_client.get("/api/worklog?part=insights").get_json()["items"]
    dumped = json.dumps(logs, ensure_ascii=False) + json.dumps(st, ensure_ascii=False)
    assert "serp-test-key" not in dumped and "fc-test-key" not in dumped and "AIza-test-key" not in dumped
    kinds = {(i["kind"], i["action"]) for i in logs}
    assert ("수집", "Amazon 진열대 · 리뷰 수집") in kinds and ("수집", "Weee! 진열대 수집") in kinds
    assert ("승인 · 선정", "이번 주 액션 선정 (규칙)") in kinds
    # 광고 제외 · soup 제외 · 같은 ASIN은 한 번
    amazon_calls = [c for c in fake_net if c[0] == "GET" and c[2].get("engine") == "amazon"]
    assert [c[2]["k"] for c in amazon_calls] == ["hangover", "korean hangover"]
    trend_calls = [c[2]["q"] for c in fake_net if c[0] == "GET" and c[2].get("engine") == "google_trends"]
    assert trend_calls == ["hangover", "hangover drink"]  # 키워드마다 따로

    shelf = auth_client.get("/api/insights/data/shelf").get_json()
    titles = [x["title"] for x in shelf["items"]]
    assert not any("Sponsored" in t or "soup" in t.lower() for t in titles)
    keys = [x["key"] for x in shelf["items"]]
    assert keys.count("amazon:B0COMP0002") == 1 and "amazon:B0OWNASIN1" in keys
    jelly = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0002")
    assert jelly["servings"] == 20 and jelly["perServing"] == 2.0 and jelly["bucket"] == "low"
    own = next(x for x in shelf["items"] if x["key"] == "amazon:B0OWNASIN1")
    assert own["own"] and own["in_stock"] is False and own["tracked"] and own["servings"] == 5
    weee = [x for x in shelf["items"] if x["source"] == "weee"]
    assert {x["title"] for x in weee} == {"Condition Hangover Drink 100ml x 10", "Aldicom Hangover Jelly 1 pack*5 pack"}
    wo = next(x for x in weee if x["own"])
    assert wo["price"] == 12.99 and wo["in_stock"] is False
    flyby = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0001")
    assert flyby["brand"] == "Flyby" and flyby["detailed"] and flyby["in_stock"] is True

    week = auth_client.get("/api/insights/weekly").get_json()
    rules_hit = {a["rule"]: a for a in week["actions"]}
    assert {"own_oos_season", "own_complaint", "season_search", "top_question"} <= set(rules_hit)
    assert "price_cut" not in rules_hit and week["shelf"]["compare"] is False  # 첫 주는 비교 안 함
    assert rules_hit["top_question"]["facts"]["질문 주제"] == "어디서 사나요 (구매처)"
    assert all(a["ai"] is False for a in week["actions"]) and week["aiReady"] is False
    assert week["shelf"]["ownWeee"]["perServing"] == 2.6

    aud = auth_client.get("/api/insights/data/audience").get_json()
    assert aud["process"]["found"] == 2 and aud["process"]["kept"] == 1
    assert aud["questions"][0]["key"] == "buy" and aud["questions"][0]["count"] == 2
    assert aud["creators"][0]["channel"] == "K-Snack Lab"
    assert all("great video" not in c["text"] for c in aud["comments"])

    rv = auth_client.get("/api/insights/data/reviews").get_json()
    own_row = next(b for b in rv["table"] if b["own"])
    assert own_row["topics"]["맛·향"]["neg"] == 2
    assert rv["summaries"][0]["text"] == "Customers like the convenience."

    tr = auth_client.get("/api/insights/trends").get_json()
    assert [k["keyword"] for k in tr["keywords"]] == ["hangover", "hangover drink"]
    assert len(tr["keywords"][0]["points"]) == 52 and tr["updatedAt"]
    nxt = auth_client.get("/api/insights/data/trends").get_json()["next"][0]
    assert nxt["name"] == "Halloween" and nxt["opportunity"] is True

    comp = auth_client.get("/api/insights/data/competitors").get_json()
    brands = {x["brand"]: x for col in comp["columns"].values() for x in col}
    assert "Flyby" in brands and brands["Flyby"]["tier"] == "category"
    assert comp["ownFound"] is True


def test_update_bad_key_fails_without_leaking(auth_client, monkeypatch, fake_net, fixed_today, sync_jobs):
    monkeypatch.setenv("SERPAPI_KEY", "bad-serp-key-000")
    assert auth_client.post("/api/insights/update", json={"items": ["amazon", "trends"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "failed"
    assert "키를 거절" in run["results"]["amazon"]["message"]
    logs = auth_client.get("/api/worklog?part=insights").get_json()["items"]
    assert logs[0]["status"] == "실패"
    assert "bad-serp-key" not in json.dumps(logs, ensure_ascii=False) + json.dumps(run, ensure_ascii=False)


def test_update_refuses_while_running(auth_client, monkeypatch):
    monkeypatch.setattr(update.jobs, "is_running", lambda name: True)
    res = auth_client.post("/api/insights/update", json={"items": ["reviews"]})
    assert res.status_code == 409


# ───────── 2주차: 가격 변동 · 신규 · 수정 ─────────

def _seed_two_weeks(app):
    from server.extensions import db
    from server.insights.models import ShelfItem
    with app.app_context():
        for week, price in (("2026-W42", 30.0), ("2026-W43", 24.0)):
            db.session.add(ShelfItem(week=week, source="amazon", key="amazon:B0COMP0001", position=1,
                                     title="Flyby Hangover Recovery Aid, 30 Count", price=price, servings=30, in_stock=True))
        db.session.add(ShelfItem(week="2026-W43", source="amazon", key="amazon:B0NEW00001", position=2,
                                 title="Brand New Hangover Patch (Pack of 6)", price=18.0, servings=6, in_stock=True))
        db.session.add(ShelfItem(week="2026-W43", source="amazon", key="amazon:B0NOSERVE1", position=3,
                                 title="Mystery hangover powder", price=9.0, in_stock=None))
        db.session.commit()


def test_second_week_changes_overrides_and_edits(app, auth_client, fixed_today):
    _seed_two_weeks(app)
    week = auth_client.get("/api/insights/weekly").get_json()
    by = {a["rule"]: a for a in week["actions"]}
    assert by["price_cut"]["grade"] == "긴급" and "$30.00 → $24.00" in by["price_cut"]["fact"]
    assert by["new_entry"]["grade"] == "참고" and week["shelf"]["compare"] is True
    assert week["shelf"]["priceChanges"] == 1 and week["shelf"]["newEntries"] == 2
    shelf = auth_client.get("/api/insights/data/shelf").get_json()
    mystery = next(x for x in shelf["items"] if x["key"] == "amazon:B0NOSERVE1")
    assert mystery["servingsSource"] == "default" and mystery["perServing"] == 9.0  # 못 읽으면 1회분으로 (확인 필요)
    flyby = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0001")
    assert flyby["change"]["pricePct"] == -20.0
    # 회분 수 직접 수정 → 모든 주차에 적용
    res = auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": 3})
    assert res.status_code == 200
    mystery = next(x for x in auth_client.get("/api/insights/data/shelf").get_json()["items"] if x["key"] == "amazon:B0NOSERVE1")
    assert mystery["servingsSource"] == "manual" and mystery["perServing"] == 3.0 and mystery["bucket"] == "mid"
    assert auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": 0}).status_code == 400
    assert auth_client.put("/api/insights/servings", json={"key": "nope", "servings": 2}).status_code == 404
    assert auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": None}).get_json()["servings"] is None
    # 한 줄 고치기: 가격 · 재고 · 자사 표시
    res = auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={"price": "10.5", "in_stock": False, "own": True})
    out = res.get_json()
    assert res.status_code == 200 and out["price"] == 10.5 and out["in_stock"] is False and out["own"] is True
    assert out["editedBy"] == "테스터"
    assert auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={"price": "abc"}).status_code == 400
    assert auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={}).status_code == 400
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["metrics"]["priceChanges"] == 1 and m["metrics"]["ownOosWeeks"] == 1


def test_watch_toggle_marks_competitor(app, auth_client, fixed_today):
    _seed_two_weeks(app)
    assert auth_client.put("/api/insights/watch", json={"brand": "Flyby", "watch": True}).get_json()["watch"] == ["Flyby"]
    comp = auth_client.get("/api/insights/data/competitors").get_json()
    flyby = next(x for col in comp["columns"].values() for x in col if x["brand"] == "Flyby")
    assert flyby["watched"] is True
    log = auth_client.get("/api/worklog?part=insights").get_json()["items"][0]
    assert log["kind"] == "승인 · 선정" and log["approval"] == "테스터"
    assert auth_client.put("/api/insights/watch", json={"brand": "Flyby", "watch": False}).get_json()["watch"] == []
    assert auth_client.put("/api/insights/watch", json={"brand": ""}).status_code == 400


# ───────── CSV 올리기 ─────────

def test_weee_review_csv_upload(auth_client):
    csv_text = "제품,브랜드,별점,날짜,제목,내용,자사\n" \
               "Aldicom Jelly,Aldicom,2,2026-10-01,,The taste is too sweet,Y\n" \
               "Condition Drink,Condition,5,2026-10-02,Great,Works great and easy to carry,N\n"
    res = auth_client.post("/api/insights/reviews/upload", json={"csv": csv_text})
    assert res.status_code == 200 and res.get_json()["added"] == 2
    assert auth_client.post("/api/insights/reviews/upload", json={"csv": csv_text}).get_json()["added"] == 0  # 같은 리뷰는 한 번
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    own = next(b for b in rv["table"] if b["own"])
    assert own["topics"]["맛·향"]["neg"] == 1
    cond = next(b for b in rv["table"] if b["brand"] == "Condition")
    assert cond["topics"]["효과"]["pos"] == 1 and cond["topics"]["복용 편의"]["pos"] == 1
    bad = auth_client.post("/api/insights/reviews/upload", json={"csv": "product,rating,text\nx,9,hello\n"})
    assert bad.status_code == 400 and "별점" in bad.get_json()["error"]
    assert auth_client.post("/api/insights/reviews/upload", json={"csv": "a,b\n1,2\n"}).status_code == 400
    tpl = auth_client.get("/api/insights/templates/weee-reviews.csv")
    assert tpl.status_code == 200 and "product,brand,rating" in tpl.get_data(as_text=True)


def test_trends_csv_upload(auth_client):
    csv_text = "Category: All categories\n\nWeek,hangover: (United States),hangover drink: (United States)\n" \
               "2025-10-26,45,<1\n2025-11-02,90,3\n2025-11-09,40,2\n"
    res = auth_client.post("/api/insights/trends/upload", json={"csv": csv_text})
    assert res.status_code == 200 and res.get_json()["keywords"] == ["hangover", "hangover drink"]
    tr = auth_client.get("/api/insights/trends").get_json()
    hang = next(k for k in tr["keywords"] if k["keyword"] == "hangover")
    assert hang["points"][1] == {"date": "2025-11-02", "value": 90.0}
    drink = next(k for k in tr["keywords"] if k["keyword"] == "hangover drink")
    assert drink["points"][0]["value"] == 0.0
    assert auth_client.post("/api/insights/trends/upload", json={"csv": "hello,world\n1,2\n"}).status_code == 400


# ───────── AI (가짜 OpenAI) ─────────

def test_ai_rewrite_and_monthly_with_mocked_openai(app, auth_client, monkeypatch, fixed_today, sync_jobs):
    from server.insights import ai as insights_ai
    _seed_two_weeks(app)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-key-zzzz")
    seen = []

    def fake_openai(system, user, *, schema=None, schema_name="result", **kw):
        seen.append(schema_name)
        if schema_name == "weekly_actions":
            return {"items": [{"id": a["id"], "fact": "AI 사실 " + a["id"], "why": "AI 이유", "todo": "AI 할 일"}
                              for a in user["actions"]]}
        if schema_name == "review_topics":
            return {"items": [{"id": r["id"], "topics": ["가격"], "sentiment": "-", "summary_ko": "비싸요"} for r in user["reviews"]]}
        if schema_name == "monthly_review":
            return {"summary": ["한 줄", "두 줄", "세 줄", "네 줄"], "insights": [], "competitor_trends": ["Flyby 가격 인하"],
                    "changes": [], "personas": [{"name": "파티 다음 날 직장인", "oneLine": "아침 회의가 걱정", "needs": "간편함",
                                                 "message": "가볍게 챙기는 리추얼"}] * 4,
                    "swot": {"S": ["a"], "W": ["b"], "O": ["c"], "T": ["d"]}, "target": "21~34세", "positioning": "간편한 리추얼"}
        raise AssertionError(schema_name)

    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    auth_client.post("/api/insights/reviews/upload", json={"csv": "product,rating,text\nX,2,too pricey for me\n"})
    res = auth_client.post("/api/insights/update", json={"items": ["reviews", "actions", "monthly"]})
    assert res.status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done", run["results"]
    assert seen == ["review_topics", "weekly_actions", "monthly_review"]
    week = auth_client.get("/api/insights/weekly").get_json()
    assert week["aiUsed"] and all(a["ai"] and a["fact"].startswith("AI 사실") for a in week["actions"])
    assert week["actions"][0]["base"]["fact"].startswith("경쟁 상품")  # 규칙 문장도 함께
    # 같은 데이터면 다시 부르지 않음 (저장된 문장)
    auth_client.post("/api/insights/update", json={"items": ["actions"]})
    assert seen.count("weekly_actions") == 1
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["ai"] and m["summary"] == ["한 줄", "두 줄", "세 줄"] and len(m["strategy"]["personas"]) == 3
    personas = auth_client.get("/api/insights/personas").get_json()
    assert personas[0] == {"name": "파티 다음 날 직장인", "oneLine": "아침 회의가 걱정", "needs": "간편함", "message": "가볍게 챙기는 리추얼"}
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    assert rv["stats"]["ai"] == 1 and rv["items"][0]["summaryKo"] == "비싸요"


# ───────── 내려받기 ─────────

def test_reports_and_export(app, auth_client, fixed_today):
    from docx import Document
    from openpyxl import load_workbook
    _seed_two_weeks(app)
    res = auth_client.get("/api/insights/report/weekly.docx")
    assert res.status_code == 200 and "wordprocessingml" in res.headers["Content-Type"]
    assert "reboot-insights-weekly-2026-W43.docx" in res.headers["Content-Disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(res.data)).paragraphs)
    assert "이번 주 액션" in text and "[긴급]" in text and "다가오는 시즌" in text and "Halloween" in text
    res = auth_client.get("/api/insights/report/monthly.docx")
    text = "\n".join(p.text for p in Document(io.BytesIO(res.data)).paragraphs)
    assert res.status_code == 200 and "월간 리포트" in text and "전략 업데이트가 없어요" in text
    res = auth_client.get("/api/insights/export.xlsx")
    assert res.status_code == 200 and "spreadsheetml" in res.headers["Content-Type"]
    wb = load_workbook(io.BytesIO(res.data))
    assert wb.sheetnames == ["진열대(주차별)", "구매자 리뷰", "VOC(유튜브·리뷰)", "검색 트렌드", "경쟁사 후보·점수", "이번 주 액션", "작업 로그"]
    assert wb["진열대(주차별)"].max_row == 1 + 4


# ───────── 파서 (수집기 단위) ─────────

def test_weee_markdown_parser():
    from server.insights.weee import parse_listing, parse_product_page
    items = parse_listing(WEEE_MD)
    assert [i["id"] for i in items] == ["12345", "67890", "11111"]
    assert items[0]["title"] == "Condition Hangover Drink 100ml x 10" and items[0]["price"] == 19.99 and items[0]["in_stock"] is True
    assert items[1]["in_stock"] is False and items[1]["price"] == 12.99
    page = parse_product_page("# Aldicom Hangover Jelly\n\n$12.99\n\nA refreshing jelly stick made with lemon and lime for your next good moment.\n",
                              "https://www.sayweee.com/en/product/Aldicom-Jelly/67890")
    assert page["id"] == "67890" and page["price"] == 12.99 and page["in_stock"] is True and "refreshing" in page["message"]


def test_amazon_search_parser_skips_ads():
    from server.insights.amazon import parse_search_item, stock_state
    item = parse_search_item({"asin": "b0comp0001", "title": " Flyby ", "price": "$29.99", "rating": "4.3", "reviews": "5,400"})
    assert item["asin"] == "B0COMP0001" and item["price"] == 29.99 and item["reviews"] == 5400 and item["title"] == "Flyby"
    assert stock_state("Currently unavailable.") is False and stock_state("Only 3 left in stock") is True and stock_state("") is None
