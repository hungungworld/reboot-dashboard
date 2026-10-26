"""카드뉴스 편집기 (2026-10-25) — 내용 넣기 → AI 덱 구성 → 템플릿에 얹기 → 화면에서 바로 수정 → PNG · ZIP.

- 덱(deck): 장마다 layout(cover · list · quote · stat · outro)과 문구 칸, 사진(브랜드 자산 ID 또는 올린 파일), 글자 크기 배율
- 저장: 버전 하나 = 덱 하나. 덱 JSON은 버전 snapshot["deck"], 본문은 기존 카드뉴스 형식('## S1 · 역할' + Headline/Body)으로
  써서 키워드 확인 · 문구별 컴플라이언스 · ZIP 내보내기가 그대로 동작
- 숫자: stat 카드의 숫자가 넣은 자료 · 제품 사실에 없으면 numberCheck로 표시 (지어낸 숫자를 그대로 두지 않음)
- 첨부: .md · .txt · .csv는 글자 그대로, .docx는 python-docx로 문단 · 표 글자만 뽑음 (PDF는 아직 미지원)
- AI 배경 (2026-10-25): 카드마다 이미지 AI 1번 — 우리 제품 · Rebu 사진(data/assets)을 참고 이미지로 넣고 글자 없이 배경만 그림.
  버튼을 눌렀을 때만 실행 · 뒤에서 실행(content_jobs, kind=cardnews_bg) · 시간 초과는 자동 재제출하지 않음 (media.py와 같은 규칙).
  어느 카드에 배경을 넣을지는 덱을 만들 때 AI가 판단(bgAi · bgReason)하거나, 장 수를 넣으면 AI가 다시 고르거나, 사람이 직접 고름
"""
import io
import re
import secrets
import time

from flask import current_app

from server import jobs, worklog
from server.core.ai import AIError
from server.extensions import db
from server.models import iso, utcnow
from . import assets, brandkit, media, prompts, providers, repository, review, schemas
from .models import ContentJob, ContentVersion

LAYOUTS = ("cover", "list", "quote", "stat", "outro")
LAYOUT_LABELS = {"cover": "표지", "list": "목록", "quote": "인용", "stat": "숫자", "outro": "마무리"}
TEMPLATES = ("neon", "clean", "photo")    # 화면 data.js cardnews.templates와 같은 ID (CSS .cnt-cn-t-<ID>)
RATIOS = ("1:1", "4:5")
TONES = {"info": "정보형 — 차분하고 정확하게", "casual": "캐주얼 — 친구에게 말하듯 가볍게",
         "hook": "후킹형 — 첫 장에서 멈추게 하는 강한 문장"}
LANGUAGES = {"en": "미국 영어 (US English)", "ko": "한국어"}
COUNTS = (0, 5, 7, 9)                     # 0 = AUTO (자료 양을 보고 5~9장)
SIZE_KEYS = ("kicker", "title", "sub", "items", "quote", "number", "desc", "cta")
SOURCE_MAX = 12000                        # AI에 보내는 자료 글자 수
EXTRACT_MAX_BYTES = 5 * 1024 * 1024
PHOTO_MAX_BYTES = 6 * 1024 * 1024
TEXT_EXTS = (".md", ".markdown", ".txt", ".csv")
RETRY_WAITS = (1.5, 4.0)                  # 붐빔(429 · 502)일 때만 기다렸다 다시 (시간 초과 504는 이미 오래 기다려서 제외)
RETRY_STATUSES = (429, 502)
PHOTO_RE = re.compile(r"^c(\d+)-cn(bg)?-[0-9a-f]{16}\.(jpg|png|webp)$")   # cn = 올린 사진 · cnbg = AI 배경
CARD_ID_RE = re.compile(r"^[a-z0-9]{4,16}$")
NUM_RE = re.compile(r"\d[\d,.]*")
BG_KIND = "cardnews_bg"
BG_MAX = 12            # 한 번에 만들 수 있는 배경 수 (덱 최대 장수와 같음)
BG_REFS = 6            # 배경 한 장에 보낼 참고 사진 수 (선택 제품 번갈아 + Rebu)


def _clip(v, n):
    return str(v if v is not None else "").strip()[:n]


# ───────── 덱 검증 ─────────

def _photo(p, campaign_id):
    """{asset: 자산 ID} · {file: 이 캠페인 파일, ai: AI 배경이면 True} · None. full=True면 카드 전체 배경으로 깖."""
    if not isinstance(p, dict):
        return None
    full = bool(p.get("full"))
    if p.get("asset") and assets.get(str(p["asset"])):
        return {"asset": str(p["asset"]), "full": full}
    m = PHOTO_RE.match(str(p.get("file") or ""))
    if m and int(m.group(1)) == campaign_id:
        ai = bool(m.group(2))
        return {"file": p["file"], "full": full or ai, "ai": ai}
    return None


def _sizes(s):
    out = {}
    for k, v in (s or {}).items() if isinstance(s, dict) else ():
        if k in SIZE_KEYS:
            try:
                f = round(float(v), 2)
            except (TypeError, ValueError):
                continue
            if 0.5 <= f <= 2.0 and f != 1.0:
                out[k] = f
    return out


def norm_card(c, campaign_id):
    c = c if isinstance(c, dict) else {}
    layout = c.get("layout") if c.get("layout") in LAYOUTS else "list"
    items = []
    for it in c.get("items") or []:
        if isinstance(it, dict) and (_clip(it.get("head"), 120) or _clip(it.get("body"), 400)):
            items.append({"head": _clip(it.get("head"), 120), "body": _clip(it.get("body"), 400)})
    cid = str(c.get("id") or "")
    return {"id": cid if CARD_ID_RE.match(cid) else secrets.token_hex(4),
            "visual": _clip(c.get("visual"), 400), "bgAi": bool(c.get("bgAi")), "bgReason": _clip(c.get("bgReason"), 160),
            "layout": layout, "kicker": _clip(c.get("kicker"), 80), "title": _clip(c.get("title"), 160),
            "sub": _clip(c.get("sub"), 300), "items": items[:5], "quote": _clip(c.get("quote"), 300),
            "number": _clip(c.get("number"), 24), "desc": _clip(c.get("desc"), 300), "cta": _clip(c.get("cta"), 120),
            "photo": _photo(c.get("photo"), campaign_id), "sizes": _sizes(c.get("sizes")),
            "numberCheck": bool(c.get("numberCheck")) and layout == "stat"}


def norm_deck(d, campaign_id):
    """화면 · AI에서 온 덱을 저장 가능한 모양으로. 장이 없으면 ValueError."""
    d = d if isinstance(d, dict) else {}
    cards = [norm_card(c, campaign_id) for c in (d.get("cards") or [])[:12] if isinstance(c, dict)]
    if not cards:
        raise ValueError("카드가 한 장도 없어요.")
    tags = []
    for t in d.get("hashtags") or []:
        t = _clip(t, 40).replace(" ", "")
        if t:
            tags.append(t if t.startswith("#") else "#" + t)
    return {"deckTitle": _clip(d.get("deckTitle"), 120) or "카드뉴스",
            "template": d.get("template") if d.get("template") in TEMPLATES else TEMPLATES[0],
            "ratio": d.get("ratio") if d.get("ratio") in RATIOS else "4:5",
            "language": d.get("language") if d.get("language") in LANGUAGES else "en",
            "tone": d.get("tone") if d.get("tone") in TONES else "info",
            "brand": _clip(d.get("brand"), 40) or "REBOOT",   # 공식 계정명은 미제공 — 지어내지 않고 브랜드명만
            "caption": _clip(d.get("caption"), 2200), "hashtags": tags[:12],
            "needsCheck": [_clip(x, 200) for x in d.get("needsCheck") or [] if isinstance(x, str) and x.strip()][:12],
            "demo": bool(d.get("demo")), "bgPlanned": bool(d.get("bgPlanned")), "cards": cards}


def fix_structure(cards):
    """첫 장은 cover, 마지막 장은 outro (AI가 어겨도 화면 구조는 지킴)."""
    if cards:
        cards[0]["layout"] = "cover"
    if len(cards) > 1:
        cards[-1]["layout"] = "outro"
    return cards


def _plain(text):
    return (text or "").replace("**", "")


def deck_body(deck):
    """덱 → 기존 카드뉴스 본문 형식 (키워드 확인 · 문구별 컴플라이언스 · 내보내기용)."""
    out = [f"# {_plain(deck['deckTitle'])}", ""]
    for i, c in enumerate(deck["cards"], 1):
        out.append(f"## S{i} · {LAYOUT_LABELS[c['layout']]}")
        if c["kicker"]:
            out.append(f"Kicker: {_plain(c['kicker'])}")
        out.append(f"Headline: {_plain(c['title'])}")
        body = [c["sub"], c["quote"] and f"\"{c['quote']}\"", c["number"] and f"{c['number']} — {c['desc']}",
                c["layout"] != "stat" and c["desc"]]
        body += [f"{it['head']}: {it['body']}" if it["head"] else it["body"] for it in c["items"]]
        for b in body:
            if b:
                out.append(f"Body: {_plain(b)}")
        if c["cta"]:
            out.append(f"CTA: {_plain(c['cta'])}")
        if c["photo"]:
            what = c["photo"].get("asset") or ("AI 배경" if c["photo"].get("ai") else "올린 사진")
            out.append(f"Visual note: (asset: {what})")
        out.append("")
    if deck["caption"]:
        out += ["## Caption", _plain(deck["caption"]), ""]
    if deck["hashtags"]:
        out += ["## Hashtags", " ".join(deck["hashtags"]), ""]
    return "\n".join(out).rstrip() + "\n"


def mark_numbers(cards, source, facts):
    """stat 카드의 숫자가 자료 · 제품 사실에 없으면 numberCheck=True (화면에 '숫자 확인 필요')."""
    known = set(NUM_RE.findall(source or "")) | set(NUM_RE.findall(prompts.dumps(facts.get("products") or [])))
    known = {k.replace(",", "").rstrip(".") for k in known}
    for c in cards:
        if c["layout"] != "stat":
            continue
        nums = {n.replace(",", "").rstrip(".") for n in NUM_RE.findall(c["number"])}
        c["numberCheck"] = not nums or not nums <= known
    return cards


# ───────── AI 생성 ─────────

def build_user(*, source, opts, sheet, decisions, facts, labels):
    count = opts["count"]
    payload = {
        "task": (f"카드뉴스 {count}장" if count else "카드뉴스 5~9장 (자료 양을 보고 정함)") + "을 구성한다.",
        "language": LANGUAGES[opts["language"]], "tone": TONES[opts["tone"]], "ratio": opts["ratio"],
        "sheet": {"channels": labels["channels"], "target": labels["target"], "brand_focus": sheet["brand_focus"],
                  "goal": sheet["goal"] or "미정", "extra_request": sheet["extra_request"] or ""},
        "decisions": {k: v.get("value") for k, v in decisions.items() if k in ("slogan_choice", "use_character")},
        "brand_facts": brandkit.brand_view(facts), "product_facts": facts.get("products") or [],
        "character_use": assets.character_choice(decisions),
        "brand_assets": assets.catalog_for(sheet["product_ids"], assets.character_choice(decisions),
                                           assets.product_flavors_of(decisions, sheet["product_ids"],
                                                                     repository.products_index())),
    }
    return (prompts.dumps(payload) + "\n\n<user_material>\n" + source[:SOURCE_MAX] + "\n</user_material>\n\n"
            + prompts.EVIDENCE_REMINDER)


def _call(system, user):
    for i in range(len(RETRY_WAITS) + 1):
        try:
            return providers.call_json(system, user, schema=schemas.CARDNEWS, name="content_cardnews",
                                       action="카드뉴스 덱 구성", max_tokens=6000, timeout=90)
        except AIError as e:
            if e.status not in RETRY_STATUSES or i == len(RETRY_WAITS):
                if e.status in RETRY_STATUSES:
                    raise AIError(f"{e} (기다렸다 {len(RETRY_WAITS)}번 더 시도했지만 안 됐어요. 1~2분 뒤 다시 눌러 주세요.)", e.status)
                raise
            time.sleep(RETRY_WAITS[i])


def generate(source, opts, *, campaign_id, sheet, decisions, facts, labels):
    """→ 덱 (저장 전). AI 실패는 AIError."""
    system = prompts.cardnews_system(LANGUAGES[opts["language"]])
    out = _call(system, build_user(source=source, opts=opts, sheet=sheet, decisions=decisions, facts=facts,
                                   labels=labels))
    cards = []
    for c in out.get("cards") or []:
        if not isinstance(c, dict):
            continue
        c = dict(c, photo={"asset": c.get("photo_asset")} if c.get("photo_asset") else None, visual=c.get("visual_en"),
                 bgAi=c.get("bg_ai"), bgReason=c.get("bg_reason_ko"))
        cards.append(c)
    if not cards:
        raise AIError("AI가 카드를 돌려주지 않았어요. 다시 요청하세요.")
    deck = norm_deck({"deckTitle": out.get("deckTitle"), "caption": out.get("caption"), "hashtags": out.get("hashtags"),
                      "needsCheck": out.get("needs_check"), "cards": cards, "template": opts["template"],
                      "ratio": opts["ratio"], "language": opts["language"], "tone": opts["tone"], "bgPlanned": True},
                     campaign_id)
    fix_structure(deck["cards"])
    mark_numbers(deck["cards"], source, facts)
    return deck, _clip(out.get("note_ko"), 300)


def check(deck, sheet, facts):
    return review.keyword_check(deck_body(deck), "cardnews", sheet, facts)


# ───────── 첨부 파일 → 글자 ─────────

def _decode(data):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _docx_text(data):
    import docx
    doc = docx.Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_text(filename, data):
    """→ 글자. 지원하지 않는 형식 · 읽기 실패는 ValueError (화면에 그대로 보여 줄 문장)."""
    name = (filename or "").lower()
    if name.endswith(TEXT_EXTS):
        text = _decode(data)
    elif name.endswith(".docx"):
        if data[:2] != b"PK":
            raise ValueError("Word(.docx) 파일이 아니에요.")
        try:
            text = _docx_text(data)
        except Exception:  # 깨진 파일 · 암호 걸린 파일 등 — 원인 문구는 화면에 내보내지 않음
            raise ValueError("Word 파일을 읽지 못했어요. 내용을 복사해 직접 입력해 주세요.")
    elif name.endswith(".pdf"):
        raise ValueError("PDF는 아직 지원하지 않아요. 내용을 복사해 직접 입력하거나 .docx · .md로 올려 주세요.")
    else:
        raise ValueError("직접 입력 또는 .md · .txt · .csv · .docx 파일을 올려 주세요.")
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n")).strip()
    if not text:
        raise ValueError("파일에서 글자를 찾지 못했어요.")
    return text


def photo_name(campaign_id):
    return f"c{campaign_id}-cn-{secrets.token_hex(8)}.jpg"


# ───────── AI 배경 — 어느 카드에 넣을지 추천 ─────────

def card_text(c):
    bits = [c.get("kicker"), c.get("title"), c.get("sub"), c.get("quote"), c.get("number"), c.get("desc"), c.get("cta")]
    bits += [f"{it.get('head', '')} {it.get('body', '')}" for it in c.get("items") or []]
    return _plain(" / ".join(b for b in bits if b and str(b).strip()))[:400]


def plan_backgrounds(deck, count, *, sheet, decisions, facts):
    """→ [{id, use, visual, reason}] (카드 순서대로). count=0이면 AI가 장 수까지 판단, 1 이상이면 정확히 그 장 수."""
    cards = deck["cards"]
    payload = {
        "task": (f"AI 배경을 넣을 카드를 정확히 {count}장 고른다." if count else "AI 배경이 어울리는 카드를 내용에 따라 고른다 (장 수도 판단)."),
        "cards": [{"index": i + 1, "layout": c["layout"], "text": card_text(c), "visual_en": c["visual"]} for i, c in enumerate(cards)],
        "brand_facts": brandkit.brand_view(facts), "product_facts": facts.get("products") or [],
        "character_use": assets.character_choice(decisions), "goal": sheet.get("goal") or "미정",
    }
    out = providers.call_json(prompts.CARDNEWS_BG_PLAN_V1, prompts.dumps(payload), schema=schemas.CARDNEWS_BG_PLAN,
                              name="content_cardnews_bg_plan", action="카드뉴스 AI 배경 추천", max_tokens=3000, timeout=60)
    got = {}
    for x in out.get("cards") or []:
        try:
            i = int(x.get("index")) - 1
        except (TypeError, ValueError):
            continue
        if 0 <= i < len(cards) and i not in got:
            got[i] = {"use": bool(x.get("use")), "visual": _clip(x.get("visual_en"), 400), "reason": _clip(x.get("reason_ko"), 160)}
    plan = [{"id": c["id"], "use": got.get(i, {}).get("use", False), "visual": got.get(i, {}).get("visual") or c["visual"],
             "reason": got.get(i, {}).get("reason", "")} for i, c in enumerate(cards)]
    if count:  # 장 수를 정했으면 정확히 맞춤 — 넘치면 뒤에서 빼고, 모자라면 표지 → 앞쪽 카드부터 채움
        picked = [p for p in plan if p["use"]][:count]
        for p in plan:
            if len(picked) >= count:
                break
            if p not in picked:
                p["reason"] = p["reason"] or "요청한 장 수를 채우려고 고름"
                picked.append(p)
        for p in plan:
            p["use"] = p in picked
    return plan


# ───────── AI 배경 — 작업 (content_jobs · 뒤에서 실행) ─────────

BG_ZONE = {  # 템플릿마다 글자가 올라가는 자리를 비워 두게 함
    "neon": "Keep the top 45% of the frame dark, calm and uncluttered (deep black or a soft dark gradient); "
            "a designer will place white headline text there. Put the main subject in the lower half.",
    "clean": "Keep the top 45% of the frame bright, calm and uncluttered (white or a very light surface); "
             "a designer will place dark headline text there. Put the main subject in the lower half.",
    "photo": "Keep the lower 45% of the frame simple and uncluttered; a white text box will cover it. "
             "Put the main subject in the upper half.",
}
BG_MOOD = {"neon": "dark, moody studio with neon lime (#E4FA18) accent lighting", "clean": "bright, airy, clean white studio",
           "photo": "natural, vivid lifestyle photography"}
BG_SAFE = ("No alcohol, no drinking scenes, no bars or bottles of liquor, no medicine or pills, no people who look sick or tired. "
           "No guns or revolver shapes.")


def bg_dict(j):
    s = j.snapshot or {}
    return {"id": j.id, "cardId": s.get("card_id"), "status": j.status, "error": j.error, "file": s.get("file"),
            "url": repository.file_url(s.get("file")) if s.get("file") else None, "visual": s.get("visual", ""),
            "createdAt": iso(j.created_at)}


def bg_jobs(campaign_id, limit=60):
    repository.settle_stale_jobs(campaign_id)
    return (ContentJob.query.filter_by(campaign_id=campaign_id, kind=BG_KIND)
            .order_by(ContentJob.created_at.desc()).limit(limit).all())


def default_visual(c, deck):
    """장면 설명이 없는 카드(예시 · 그대로 나누기 · 직접 추가)의 기본 장면."""
    return (f"A stylish product scene that matches this slide message: \"{card_text(c)[:160]}\". "
            "Reboot jelly pouches or bottles as the hero, with props that fit the message.")


def start_backgrounds(c, items, deck, *, user):
    """items: [{id, visual}] → 작업 목록. 카드마다 작업 하나 (실패해도 다른 카드는 그대로)."""
    sheet = repository.sheet_of(c)
    decisions = c.decisions
    per_product = assets.product_flavors_of(decisions, sheet["product_ids"], repository.products_index())
    choice = assets.character_choice(decisions)
    model, size = providers.image_model(), media.image_size(deck["ratio"])
    out = []
    for it in items:
        visual = _clip(it.get("visual"), 400)
        card = next((x for x in deck["cards"] if x["id"] == it.get("id")), None)
        if card is None:
            continue
        visual = visual or default_visual(card, deck)
        character = choice == "사용" or (choice == "AI 판단" and "rebu" in visual.lower())
        job = ContentJob(id=secrets.token_hex(16), campaign_id=c.id, kind=BG_KIND, status="queued", model=model, size=size,
                         instructions=visual, created_by=user)
        job.snapshot = {"provider": "openai", "card_id": card["id"], "visual": visual, "template": deck["template"],
                        "ratio": deck["ratio"], "layout": card["layout"], "product_ids": sheet["product_ids"],
                        "per_product": per_product, "character": character, "ref_skip": assets.ref_skip_of(decisions)}
        db.session.add(job)
        out.append(job)
    db.session.commit()
    return out


def launch(app, job_id):
    jobs.run_background(app, f"content-cnbg-{job_id}", run_bg_job, job_id)


def bg_prompt(snap, refs, facts):
    tpl = snap.get("template") if snap.get("template") in TEMPLATES else "neon"
    head = (f"Background image for one slide of an Instagram card-news carousel ({snap.get('ratio') or '4:5'}) for Reboot, "
            f"a US jelly and drink brand. Scene: {snap.get('visual')} Style: premium product photography, {BG_MOOD[tpl]}. "
            f"{BG_ZONE[tpl]} {BG_SAFE}")
    if refs:
        return head + "\n\n" + media._ref_lines(refs, facts, 1) + "\n" + media.REF_TEXT + media.REF_NO_TEXT
    return head + "\n\n" + media.NO_TEXT


def run_bg_job(job_id):
    job = db.session.get(ContentJob, job_id)
    if job is None or job.status != "queued":
        return  # 이미 실행했거나 멈춤 처리됨 — 다시 제출하지 않음
    job.status, job.started_at = "running", utcnow()
    db.session.commit()
    snap = dict(job.snapshot)
    started = time.time()
    try:
        pids = snap.get("product_ids") or []
        refs = (assets.pick_refs(pids, per_product=snap.get("per_product") or {}, character=bool(snap.get("character")),
                                 limit=BG_REFS, skip=snap.get("ref_skip") or ()) if providers.image_supports_edit(job.model) else [])
        job.prompt = bg_prompt(snap, refs, brandkit.product_facts(pids))
        db.session.commit()  # 프롬프트가 있으면 '제공사에 보냄'으로 봄 (멈추면 자동 재제출 없이 결과 확인 불가)
        if refs:
            png, usage = providers.edit_image([str(assets.path_of(a)) for a in refs], job.prompt, job.size, job.model)
        else:
            png, usage = providers.generate_image(job.prompt, job.size, job.model)
    except providers.SubmissionUnknown as e:
        return _bg_fail(job, str(e), "submission_unknown")
    except providers.ProviderError as e:
        return _bg_fail(job, str(e))
    except Exception:  # 원문은 서버 로그에만
        current_app.logger.exception("카드뉴스 AI 배경 작업 실패")
        return _bg_fail(job, "배경 생성 중 알 수 없는 오류가 났어요.")
    job.status = "saving"
    db.session.commit()
    ext = media.file_ext(png)
    if ext not in (".png", ".jpg", ".webp"):
        return _bg_fail(job, "배경 결과 파일 형식을 알 수 없어 저장하지 않았어요.")
    name = f"c{job.campaign_id}-cnbg-{secrets.token_hex(8)}{ext}"
    try:
        (media.upload_dir() / name).write_bytes(png)
    except OSError:
        return _bg_fail(job, "배경은 생성됐지만 파일을 저장하지 못했어요 (서버 저장 공간 확인).")
    job.snapshot = dict(snap, file=name, references=[a["id"] for a in refs], usage=usage, seconds=round(time.time() - started, 1))
    job.status, job.finished_at = "succeeded", utcnow()
    db.session.commit()
    worklog.add("content", "AI 프롬프트", "카드뉴스 AI 배경 생성",
                detail={"model": job.model, "size": job.size, "prompt": job.prompt[:1500], "campaign": job.campaign_id,
                        "references": len(refs), "usage": usage}, user_name=job.created_by)


def _bg_fail(job, message, status="failed"):
    job.status, job.error, job.finished_at = status, message[:400], utcnow()
    db.session.commit()
    worklog.add("content", "AI 프롬프트", "카드뉴스 AI 배경 생성", status="실패",
                detail={"model": job.model, "size": job.size, "error": message, "campaign": job.campaign_id},
                user_name=job.created_by)


# ───────── 내 캠페인 썸네일 ─────────

COVER_CARDS = 3


def deck_covers(campaigns):
    """{캠페인 ID: {template, ratio, brand, count, cards(앞 3장)}} — 최종본 덱 우선, 없으면 가장 최근 덱."""
    if not campaigns:
        return {}
    finals = {c.id: c.final_version_id for c in campaigns}
    rows = (ContentVersion.query.filter(ContentVersion.campaign_id.in_(list(finals)), ContentVersion.kind == "cardnews")
            .order_by(ContentVersion.id.desc()).all())
    out = {}
    for v in rows:
        if v.campaign_id in out and finals.get(v.campaign_id) != v.id:
            continue
        deck = v.snapshot.get("deck")
        if not deck or not deck.get("cards"):
            continue
        out[v.campaign_id] = {"template": deck.get("template"), "ratio": deck.get("ratio"), "brand": deck.get("brand"),
                              "count": len(deck["cards"]), "cards": deck["cards"][:COVER_CARDS]}
    return out
