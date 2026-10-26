"""문구 검토 요청 (2026-10-25) — 단건 심의 문구를 법무·인허가·품질팀에 보내고 회신을 받음.

흐름: 단건 심의 [검토 요청서 보내기] → 팀마다 1회용 회신 링크를 만들고, 걸린 규칙을 정리한 엑셀 요청서(workbook)를
     첨부해 Gmail로 등록된 담당자에게 보냄 (2026-10-26 · 예전 Slack 알림 대신 · Gmail은 CS 탭 연결을 같이 씀)
     → 담당자가 메일 본문의 링크(/review/<토큰> · 로그인 없음)를 열어 승인·수정 필요·반려 + 의견 → 대시보드에 팀별 상태.
     회신 알림은 Slack 웹훅(SLACK_WEBHOOK_URL)이 있을 때만 보냄.
보안: 링크 토큰은 저장하지 않고 지문(SHA-256)만 저장 · 14일 뒤 만료 · 링크는 그 팀·그 요청 하나에만 회신 가능.
규칙 판정은 화면이 보낸 값을 믿지 않고 서버가 다시 점검(server_review)해서 함께 보냄.
저장: DATA_DIR/compliance/review-requests.json (깃에 올리지 않음 · 최근 300건)
"""
import hashlib
import io
import json
import re
import secrets
import threading
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

from flask import current_app

from server.envkeys import get as env

TEAMS = {"legal": "법무팀", "regulatory": "인허가팀", "quality": "품질팀"}
DECISIONS = ("승인", "수정 필요", "반려")
EXPIRES = timedelta(days=14)
KEEP = 300
SLACK_RE = re.compile(r"^https://hooks\.slack\.com/services/[A-Za-z0-9/_-]+$")
_lock = threading.Lock()


class RequestError(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def _now():
    return datetime.now(timezone.utc)


def _iso(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def _hash(token):
    return hashlib.sha256(str(token).encode("utf-8")).hexdigest()


def _path():
    return Path(current_app.config["DATA_DIR"]) / "compliance" / "review-requests.json"


def _load():
    try:
        data = json.loads(_path().read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def _save(items):
    p = _path()
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".json.tmp")
    for attempt in range(3):   # Windows에서 파일이 잠깐 잡혀 있을 수 있음 (번역 캐시와 같은 처리)
        try:
            tmp.write_text(json.dumps(items[:KEEP], ensure_ascii=False), encoding="utf-8")
            tmp.replace(p)
            return
        except OSError:
            time.sleep(0.2 * (attempt + 1))
    raise RequestError("검토 요청을 저장하지 못했습니다. 잠시 뒤 다시 시도하세요.", 500)


def public(r):
    """화면에 보내는 모양 (토큰 지문은 빼고)."""
    return {k: v for k, v in r.items() if k != "teams"} | {
        "teams": {t: {k: v for k, v in x.items() if k != "tokenHash"} for t, x in r["teams"].items()},
        "expired": _now() > datetime.fromisoformat(r["expiresAt"].replace("Z", "+00:00"))}


# ───────── Slack ─────────

def _esc(t):
    return str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _slack(text):
    """보내면 '', 못 보내면 화면에 보여 줄 이유. 요청 자체는 Slack 실패와 상관없이 저장됨."""
    url = env("SLACK_WEBHOOK_URL").strip()
    if not url:
        return ""   # 웹훅을 안 쓰는 팀 (2026-10-26 · 요청은 메일로 보내고 Slack은 회신 알림만 선택)
    if not SLACK_RE.match(url):
        return "Slack 웹훅 주소 모양이 다릅니다 (https://hooks.slack.com/services/… )."
    req = urllib.request.Request(url, data=json.dumps({"text": text}).encode("utf-8"),
                                 headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=15) as r:
            return "" if r.status == 200 else f"Slack이 메시지를 받지 않았습니다 ({r.status})."
    except urllib.error.HTTPError as e:
        return f"Slack이 메시지를 받지 않았습니다 ({e.code})."
    except (urllib.error.URLError, TimeoutError) as e:
        return f"Slack에 연결하지 못했습니다 ({type(e).__name__})."


def _cut(t, n):
    t = str(t or "").strip()
    return t if len(t) <= n else t[:n - 1] + "…"


# ───────── 요청 · 회신 ─────────

def _details(details, codes):
    """화면이 보낸 규칙 설명(한 줄 결론·할 일·근거) 중 서버가 다시 점검해 걸린 규칙 것만 · 글자 수 제한."""
    details = details if isinstance(details, dict) else {}
    out = {}
    for c in codes:
        d = details.get(c) if isinstance(details.get(c), dict) else {}

        def lst(k, n):
            v = d.get(k)
            return [_cut(x, 300) for x in (v if isinstance(v, list) else [])[:n]]
        out[c] = {"say": _cut(d.get("say"), 200), "todo": lst("todo", 5), "reason": _cut(d.get("reason"), 600),
                  "alt": _cut(d.get("alt"), 300), "law": lst("law", 4), "legalNote": _cut(d.get("legalNote"), 400)}
    return out


def draft(text, ko, channel, product, ai_verdict, teams, memo, requester, details=None):
    """저장 전 요청 (서버 재검사 · 번호·회신 링크 없음) — 보내기 전 엑셀 미리 받기에도 씀."""
    from . import server_review
    text = str(text or "").strip()
    if not text or len(text) > 2000:
        raise RequestError("검토를 요청할 영어 문구가 없거나 너무 깁니다 (2000자 이하).")
    teams = [t for t in dict.fromkeys(teams or []) if t in TEAMS]
    if not teams:
        raise RequestError("검토할 팀을 하나 이상 고르세요 (법무 · 인허가 · 품질).")
    product = product if product in server_review.products() else "all"
    rule = server_review.review(text, "b2b" if channel == "b2b" else "cs", product)
    now = _now()
    desc = _details(details, rule["codes"])
    hits = [x | {"terms": rule["hits"].get(x["code"], [])} | desc.get(x["code"], {}) for x in rule["items"]]
    hits.sort(key=lambda x: (-server_review.LIGHT_RANK.get(x["legal"], 0), -server_review.RANK.get(x["verdict"], 0)))
    return {"id": "", "no": "", "createdAt": _iso(now), "expiresAt": _iso(now + EXPIRES),
            "text": text, "ko": _cut(ko, 2000), "channel": str(channel or "")[:20], "product": product,
            "ruleVerdict": rule["verdict"], "ruleLegal": rule["legal"], "ruleCodes": rule["codes"], "aiVerdict": str(ai_verdict or "")[:10],
            "hits": hits, "memo": _cut(memo, 1000), "requester": _cut(requester, 60),
            "teams": {t: {"label": TEAMS[t], "status": "대기", "comment": "", "by": "", "at": ""} for t in teams}, "mail": None}


def create(text, ko, channel, product, ai_verdict, teams, memo, requester, base_url, details=None, attachment=None):
    """요청 저장 → (요청, {팀: 회신 링크}). 엑셀은 file_bytes · 메일은 send_mail (2026-10-26 · Slack 알림 없앰).
    attachment: 사용자가 미리 받아 고친 요청서(XLSX bytes) — 요청번호를 넣어 이 요청의 첨부로 보관."""
    r = draft(text, ko, channel, product, ai_verdict, teams, memo, requester, details)
    edited = bool(attachment)
    if edited:
        _open_xlsx(attachment)   # 엑셀이 아니면 저장 전에 거절
    r["id"] = secrets.token_hex(6)
    r["no"] = f"CR-{r['createdAt'][2:10].replace('-', '')}-{secrets.token_hex(2).upper()}"
    links = {}
    for t, row in r["teams"].items():
        token = secrets.token_urlsafe(24)
        links[t] = f"{base_url.rstrip('/')}/review/{token}"
        row["tokenHash"] = _hash(token)
    r["edited"] = edited
    if edited:
        p = _file_path(r["id"])
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(stamp(attachment, r))
    with _lock:
        items = _load()
        items.insert(0, r)
        _save(items)
    return public(r), links


def get(rid):
    r = next((x for x in _load() if x["id"] == str(rid)), None)
    if not r:
        raise RequestError("없는 검토 요청입니다.", 404)
    return r


# ───────── 엑셀 요청서 (2026-10-26) ─────────

VERDICT_KO = {"PASS": "규칙 미탐지", "COND": "조건 확인", "REVIEW": "담당자 검토", "REVISE": "수정 필요", "BLOCK": "사용 불가"}   # 화면 LABEL과 같음
LIGHT_KO = {"red": "법 위반", "yellow": "법무 확인", "green": "팀·플랫폼 기준", "": "신호 없음"}
CHANNEL_KO = {"b2c": "B2C 광고·SNS", "amazon": "아마존 상세페이지", "cs": "CS 답변", "b2b": "B2B 바이어 자료"}   # data.js channels
PRODUCT_KO = {"jelly": "Original 젤리", "jelly-ls": "Balance 젤리", "drink": "Original 음료", "drink-ls": "Balance 음료", "all": "전체 제품"}
LIGHT_FILL = {"red": "FDE2E1", "yellow": "FFF3C4", "green": "E3F1E4"}


def _file_path(rid):
    return Path(current_app.config["DATA_DIR"]) / "compliance" / "review-files" / f"{rid}.xlsx"


def file_bytes(r):
    """이 요청의 첨부 엑셀: 사용자가 고친 파일을 올렸으면 그것, 아니면 지금 상태로 새로 만듦."""
    p = _file_path(r["id"]) if r.get("id") else None
    if p and r.get("edited") and p.exists():
        return p.read_bytes()
    return workbook(r)


def file_name(r):
    return f"검토요청서_{r.get('no') or '미리보기'}.xlsx"


def workbook(r):
    """요청서 · 걸린 표현 · 팀 회신란 3장짜리 XLSX(bytes). 회신 링크는 넣지 않음 (파일이 전달돼도 아무나 회신 못 하게)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    thin = Side(style="thin", color="C9CCC2")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="23251F")
    label_fill = PatternFill("solid", fgColor="F1F2EC")
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def title(ws, text, sub, width):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws["A1"] = text
        ws["A1"].font = Font(size=16, bold=True)
        ws["A2"] = sub
        ws["A2"].font = Font(size=10, color="6B7065")
        ws.row_dimensions[1].height = 28

    def header(ws, row, names):
        for i, n in enumerate(names, 1):
            c = ws.cell(row=row, column=i, value=n)
            c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), head_fill, center, box
        ws.row_dimensions[row].height = 22

    def landscape(ws):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1

    wb = Workbook()
    no = r.get("no") or "(보내면 붙음)"
    sub = f"요청번호 {no} · 규칙·AI 판정은 1차 참고용이며 법률 자문이 아닙니다. 최종 판단은 검토 팀이 합니다."
    hits = r.get("hits") or []

    # ① 요청서 (결재 문서처럼 항목 | 내용 두 쌍)
    ws = wb.active
    ws.title = "요청서"
    title(ws, "REBOOT 광고 문구 검토 요청서", sub, 4)
    info = [("요청번호", no, "요청일", r["createdAt"][:10]),
            ("요청자", r["requester"], "회신 기한", r["expiresAt"][:10]),
            ("사용 채널", CHANNEL_KO.get(r["channel"], r["channel"] or "—"), "제품", PRODUCT_KO.get(r["product"], r["product"] or "—")),
            ("검토 팀", " · ".join(t["label"] for t in r["teams"].values()), "걸린 규칙", f"{len(hits)}건"),
            ("규칙 판정", f"{VERDICT_KO.get(r['ruleVerdict'], r['ruleVerdict'])} ({r['ruleVerdict']})",
             "법적 신호", LIGHT_KO.get(r["ruleLegal"], r["ruleLegal"])),
            ("AI 2차 판정", VERDICT_KO.get(r["aiVerdict"], r["aiVerdict"]) if r["aiVerdict"] else "받지 않음", "", "")]
    row = 3
    for vals in info:
        row += 1
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border, cell.alignment = box, wrap
            if col % 2:
                cell.fill, cell.font = label_fill, Font(bold=True)
        if vals[2] == "법적 신호" and r["ruleLegal"] in LIGHT_FILL:
            ws.cell(row=row, column=4).fill = PatternFill("solid", fgColor=LIGHT_FILL[r["ruleLegal"]])
    for label, value, h in (("검토할 영어 문구", r["text"], 90), ("한국어 원문", r["ko"] or "—", 60), ("요청 메모", r["memo"], 48)):
        if not value:   # 메모가 없으면 줄을 만들지 않음
            continue
        row += 1
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        a = ws.cell(row=row, column=1, value=label)
        a.fill, a.font, a.alignment = label_fill, Font(bold=True), wrap
        ws.cell(row=row, column=2, value=value).alignment = wrap
        for col in (1, 2, 3, 4):
            ws.cell(row=row, column=col).border = box
        ws.row_dimensions[row].height = h
    for col, w in zip("ABCD", (16, 40, 16, 40)):
        ws.column_dimensions[col].width = w
    landscape(ws)

    # ② 걸린 표현 (위험이 큰 것부터)
    ws = wb.create_sheet("걸린 표현")
    cols = ["No", "걸린 표현", "규칙", "규칙 이름", "판정", "법적 신호", "왜 문제인가", "고치는 방법", "근거 (법령·지침)"]
    title(ws, "걸린 표현 정리", sub, len(cols))
    header(ws, 4, cols)
    for i, h in enumerate(hits, 1):
        why = "\n".join(x for x in (h.get("say"), h.get("reason")) if x)
        fix = "\n".join([f"· {t}" for t in h.get("todo") or []] + ([f"· 대체 방향: {h['alt']}"] if h.get("alt") else []))
        law = "\n".join(list(h.get("law") or []) + ([h["legalNote"]] if h.get("legalNote") else []))
        vals = [i, " · ".join(h.get("terms") or []), h["code"], h["name"],
                VERDICT_KO.get(h["verdict"], h["verdict"]), LIGHT_KO.get(h["legal"], h["legal"]), why, fix, law]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=4 + i, column=col, value=v)
            c.border, c.alignment = box, center if col in (1, 3, 5, 6) else wrap
        if h["legal"] in LIGHT_FILL:
            ws.cell(row=4 + i, column=6).fill = PatternFill("solid", fgColor=LIGHT_FILL[h["legal"]])
    if not hits:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(cols))
        ws.cell(row=5, column=1, value="규칙에 걸린 표현이 없습니다. 규칙이 놓친 위험이 있는지 검토해 주세요.").alignment = center
    for col, w in zip("ABCDEFGHI", (5, 24, 7, 22, 10, 13, 44, 44, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    landscape(ws)

    # ③ 팀 회신란 (결정은 드롭다운)
    ws = wb.create_sheet("팀 회신")
    cols = ["검토 팀", "결정", "의견 (수정 필요·반려는 이유 필수)", "검토자", "검토일"]
    title(ws, "팀 회신란", "메일의 팀별 회신 링크로 답하거나, 이 표의 결정·의견·검토자를 채워 메일로 회신해 주세요. 요청자가 대시보드에 올리면 반영됩니다.", len(cols))
    header(ws, 4, cols)
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, t in enumerate(r["teams"].values(), 5):
        vals = [t["label"], "" if t["status"] == "대기" else t["status"], t["comment"], t["by"], t["at"][:10]]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=v)
            c.border, c.alignment = box, wrap if col == 3 else center
        dv.add(f"B{i}")
        ws.row_dimensions[i].height = 48
    for col, w in zip("ABCDE", (14, 14, 60, 14, 14)):
        ws.column_dimensions[col].width = w
    landscape(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MAX_XLSX = 3 * 1024 * 1024


def _open_xlsx(content):
    from openpyxl import load_workbook
    if not content or len(content) > MAX_XLSX:
        raise RequestError("엑셀 파일이 없거나 너무 큽니다 (3MB 이하).")
    if content[:2] != b"PK":
        raise RequestError("XLSX 파일만 올릴 수 있어요.")
    try:
        return load_workbook(io.BytesIO(content))
    except Exception:
        raise RequestError("엑셀 파일을 열지 못했어요. 다른 이름으로 저장(XLSX)한 뒤 다시 올려 주세요.")


def stamp(content, r):
    """사용자가 고친 요청서에 요청번호를 넣음 (회신 엑셀을 다시 올릴 때 이 번호로 찾음)."""
    wb = _open_xlsx(content)
    if "요청서" in wb.sheetnames and str(wb["요청서"]["A4"].value or "").strip() == "요청번호":
        wb["요청서"]["B4"] = r["no"]
    for ws in wb:   # 미리보기 때 비워 둔 번호
        if isinstance(ws["A2"].value, str):
            ws["A2"] = ws["A2"].value.replace("(보내면 붙음)", r["no"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_replies(content):
    """팀이 [팀 회신] 시트를 채워 돌려준 엑셀 → 요청에 반영. (요청, 반영한 팀, 건너뛴 줄 이유)."""
    wb = _open_xlsx(content)
    if "요청서" not in wb.sheetnames or "팀 회신" not in wb.sheetnames:
        raise RequestError("검토 요청서 엑셀이 아니에요 ([요청서] · [팀 회신] 시트가 있어야 해요).")
    no = str(wb["요청서"]["B4"].value or "").strip()
    by_label = {v: k for k, v in TEAMS.items()}
    rows = []
    for row in wb["팀 회신"].iter_rows(min_row=5, values_only=True):
        if not row or not row[0]:
            continue
        rows.append([str(v).strip() if v is not None else "" for v in (list(row) + [""] * 5)[:5]])
    updated, skipped = [], []
    with _lock:
        items = _load()
        r = next((x for x in items if no and x.get("no") == no), None)
        if r is None:
            raise RequestError(f"요청번호 {no or '(없음)'}에 맞는 검토 요청을 찾지 못했어요.", 404)
        for label, decision, comment, name, _ in rows:
            t = by_label.get(label)
            if not t or t not in r["teams"]:
                skipped.append(f"{label}: 이 요청에 없는 팀")
            elif not decision:
                continue   # 아직 안 채운 팀
            elif decision not in DECISIONS:
                skipped.append(f"{label}: 결정은 승인 · 수정 필요 · 반려 중 하나")
            elif decision != "승인" and not comment:
                skipped.append(f"{label}: 수정 필요·반려는 의견이 필요해요")
            else:
                r["teams"][t].update({"status": decision, "comment": _cut(comment, 1000), "by": _cut(name, 60) or "엑셀 회신",
                                      "at": _iso(_now()), "via": "xlsx"})
                updated.append(t)
        if updated:
            _save(items)
    return public(r), updated, skipped


# ───────── 메일 (CS 탭 Gmail 연결·담당자 연락처를 같이 씀 · 등록된 담당자에게만) ─────────

TEAM_WORDS = (("legal", "법무"), ("regulatory", "인허가"), ("quality", "품질"))


def mail_setup():
    """{contacts:[{role,name,email,team}], gmail} — 받는 사람 고르기용. team = 역할 이름으로 짐작한 팀 ('' = 모름)."""
    from tabs.cs import contacts, gmail
    out = [{"role": c.get("role", ""), "name": c.get("name", ""), "email": c.get("email", ""),
            "team": next((k for k, w in TEAM_WORDS if w in c.get("role", "")), "")} for c in contacts.load()]
    return {"contacts": out, "gmail": gmail.ready()}


def send_mail(rid, to, subject, body, links, sender):
    """엑셀 요청서를 첨부해 Gmail로 보냄 → 보낸 기록. 받는 사람은 CS 담당자 연락처에 있는 주소만."""
    import base64
    from email.message import EmailMessage

    from tabs.cs import contacts, gmail
    to = list(dict.fromkeys(str(x).strip() for x in (to if isinstance(to, list) else []) if str(x).strip()))[:10]
    if not to:
        raise RequestError("받는 사람을 한 명 이상 고르세요.")
    bad = [x for x in to if not contacts.allowed(x)]
    if bad:
        raise RequestError(f"등록된 담당자에게만 보낼 수 있어요: {', '.join(bad)} (고객 응대 › 답장 설정 › 담당자 연락처)")
    subject = _cut(subject, 200)
    body = str(body or "").strip()
    if not subject or not body or len(body) > 20000:
        raise RequestError("제목과 본문을 확인해 주세요 (본문 2만 자 이하).")
    r = get(rid)
    links = links if isinstance(links, dict) else {}
    # 회신 링크는 만들 때 한 번만 화면에 보임 → 화면이 돌려보낸 링크가 이 요청 것인지 지문으로 확인
    for t, url in links.items():
        if t not in r["teams"] or not secrets.compare_digest(r["teams"][t]["tokenHash"], _hash(str(url).rsplit("/", 1)[-1])):
            raise RequestError("회신 링크가 이 요청과 맞지 않아요. 창을 닫고 다시 시도하세요.")
    if links:
        body += (f"\n\n─────────────\n팀별 회신 링크 (로그인 없음 · {r['expiresAt'][:10]}까지)\n"
                 + "\n".join(f"· {TEAMS[t]}: {links[t]}" for t in r["teams"] if t in links))
    msg = EmailMessage()
    msg["To"] = ", ".join(to)
    msg["Subject"] = subject
    msg.set_content(body)
    msg.add_attachment(file_bytes(r), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=file_name(r))
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    try:
        sent = gmail.service(interactive=False).users().messages().send(userId="me", body={"raw": raw}).execute()
    except gmail.LoginRequired as e:
        raise RequestError(str(e), 401)
    except Exception as e:
        raise RequestError(f"Gmail로 보내지 못했어요: {type(e).__name__}", 502)
    log = {"to": to, "subject": subject, "at": _iso(_now()), "by": _cut(sender, 60), "id": sent.get("id", "")}
    with _lock:
        items = _load()
        x = next((x for x in items if x["id"] == r["id"]), None)
        if x is not None:
            x["mail"] = log
            _save(items)
    return log


def find(token):
    """(요청, 팀) — 없거나 만료면 RequestError."""
    h = _hash(token)
    for r in _load():
        for t, x in r["teams"].items():
            if secrets.compare_digest(x.get("tokenHash", ""), h):
                if _now() > datetime.fromisoformat(r["expiresAt"].replace("Z", "+00:00")):
                    raise RequestError("회신 기간(14일)이 지난 링크입니다. 요청한 사람에게 다시 보내 달라고 하세요.", 410)
                return r, t
    raise RequestError("없는 링크입니다. 주소가 잘렸는지 확인하세요.", 404)


def respond(token, decision, comment, name):
    """회신 저장 (같은 링크로 다시 회신하면 덮어씀) + Slack 알림 → (요청, 팀, Slack 실패 이유)."""
    if decision not in DECISIONS:
        raise RequestError("승인 · 수정 필요 · 반려 중 하나를 고르세요.")
    name = _cut(name, 60)
    if not name:
        raise RequestError("회신하는 사람 이름을 적어 주세요.")
    comment = _cut(comment, 1000)
    if decision != "승인" and not comment:
        raise RequestError("수정 필요·반려는 이유를 적어 주세요.")
    with _lock:
        items = _load()
        r, team = find(token)
        r = next(x for x in items if x["id"] == r["id"])
        r["teams"][team].update({"status": decision, "comment": comment, "by": name, "at": _iso(_now())})
        _save(items)
    slack_error = _slack(f"*[REBOOT 문구 검토 회신]* {TEAMS[team]} *{decision}* · {_esc(name)}"
                         + (f"\n의견: {_esc(comment)}" if comment else "") + f"\n>{_esc(_cut(r['text'], 300))}")
    return r, team, slack_error


def recent(limit=50):
    return [public(r) for r in _load()[:limit]]
