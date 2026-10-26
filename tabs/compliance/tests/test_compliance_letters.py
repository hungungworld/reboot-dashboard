"""FDA 경고장 수집 · AI 페르소나 — 인터넷·키 없이 확인 (외부 요청은 모두 가짜로 바꿈)."""
import io
import json
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from tabs.compliance import ai_review, refresh, warning_letters

LETTER = """WARNING LETTER  MARCS-CMS 612345  Product: Party Gummies
Your website makes claims such as: “Party Gummies prevent hangovers so you can drink all night” and
“Protects your liver from alcohol damage”. See section 201(g)(1) of the Act, 21 U.S.C. 321(g)(1)."""


def _xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Posted Date", "Letter Issue Date", "Company Name", "Issuing Office", "Subject"])
    recent = (date.today() - timedelta(days=10)).strftime("%m/%d/%Y")
    old = (date.today() - timedelta(days=900)).strftime("%m/%d/%Y")
    ws.append([recent, recent, "Party Gummies LLC", "Human Foods Program", "Unapproved New Drugs/Misbranded"])
    ws["C2"].hyperlink = "/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/party-gummies-llc-612345-10012026"
    ws.append([recent, recent, "Pill Maker Inc", "Center for Drug Evaluation and Research", "CGMP/Finished Pharmaceuticals"])
    ws.append([old, old, "Old Juice Co", "Human Foods Program", "Juice HACCP/Adulterated"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _offline(monkeypatch):
    """자동 경로는 fda.gov 봇 차단처럼, 다른 출처는 0건으로."""
    def blocked(*a, **k):
        raise warning_letters.Blocked("fda.gov가 자동 요청을 막았습니다 (봇 차단).")
    monkeypatch.setattr(warning_letters, "_fetch", blocked)
    monkeypatch.setattr(refresh, "SOURCES", [refresh.SOURCES[0]])
    for k in ("FDA_DASHBOARD_USER", "FDA_DASHBOARD_KEY"):   # .env의 실제 키로 FDA를 부르지 않게
        monkeypatch.delenv(k, raising=False)


def test_extract_claims_skips_legal_text():
    quotes = warning_letters.extract_claims(LETTER)
    assert quotes == ["Party Gummies prevent hangovers so you can drink all night", "Protects your liver from alcohol damage"]


def test_xlsx_upload_keeps_recent_food_letters(auth_client):
    res = auth_client.post("/api/compliance/letters/xlsx", data={"file": (io.BytesIO(_xlsx()), "letters.xlsx")},
                           content_type="multipart/form-data")
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert out["letters"] == {"rows": 3, "food": 1, "added": 1}
    src = next(s for s in out["sources"] if s["id"] == "fda-warning-letters")
    assert [i["id"] for i in src["items"]] == ["612345"]
    assert src["items"][0]["url"].startswith("https://www.fda.gov/")


def test_xlsx_upload_rejects_other_files(auth_client):
    res = auth_client.post("/api/compliance/letters/xlsx", data={"file": (io.BytesIO(b"hello"), "a.xlsx")},
                           content_type="multipart/form-data")
    assert res.status_code == 400
    assert auth_client.post("/api/compliance/letters/xlsx", data={}, content_type="multipart/form-data").status_code == 400


def test_text_upload_becomes_casebook_candidates(auth_client):
    res = auth_client.post("/api/compliance/letters/text", json={
        "text": LETTER, "url": "https://www.fda.gov/x/party-gummies-llc-612345-10012026", "firm": "Party Gummies LLC", "issued": "2026-10-01"})
    assert res.status_code == 200, res.get_json()
    item = next(s for s in res.get_json()["sources"] if s["id"] == "fda-warning-letters")["items"][0]
    assert item["id"] == "612345" and len(item["quotes"]) == 2 and item["date"] == "2026-10-01"


def test_text_upload_input_checks(auth_client):
    assert auth_client.post("/api/compliance/letters/text", json={"text": "short"}).status_code == 400
    assert auth_client.post("/api/compliance/letters/text", json={"text": LETTER, "url": "https://evil.example/x"}).status_code == 400
    assert auth_client.post("/api/compliance/letters/text", json={"text": "x" * 100}).status_code == 400   # 인용 문구 없음


def test_refresh_reports_block_but_keeps_uploaded(auth_client, monkeypatch):
    _offline(monkeypatch)
    auth_client.post("/api/compliance/letters/text", json={"text": LETTER, "firm": "Party Gummies LLC"})
    out = auth_client.post("/api/compliance/monitor/refresh", json={}).get_json()
    src = next(s for s in out["sources"] if s["id"] == "fda-warning-letters")
    assert src["items"] and src["items"][0]["quotes"]
    errs = " ".join(e["error"] for e in out["errors"])
    assert "봇 차단" in errs and "Data Dashboard 키 없음" in errs


def test_ai_review_sends_persona(app, monkeypatch):
    sent = {}

    class Fake:
        def __init__(self, **kw):
            self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))

        def create(self, **kw):
            sent.update(kw)
            body = {"verdict": "BLOCK", "confidence": "high", "summary": "s", "issues": [], "rewrite": "r"}
            return SimpleNamespace(stop_reason="end_turn", model="claude-test", content=[SimpleNamespace(type="text", text=json.dumps(body))])

    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    monkeypatch.setattr(ai_review.anthropic, "Anthropic", Fake)
    with app.app_context():
        ai_review.review("Prevents hangovers.", persona="미국 광고 전문가 " + "가" * 700)
        texts = [b["text"] for b in sent["system"]]
        assert texts[0] == ai_review.SYSTEM and "never makes a verdict less strict" in texts[1]
        assert len(texts[1]) == len(ai_review.PERSONA_INTRO) + ai_review.MAX_PERSONA   # 600자에서 자름
        assert "cache_control" in sent["system"][-1]
        sent.clear()
        ai_review.review("Prevents hangovers.")
        assert len(sent["system"]) == 2   # 페르소나가 없으면 블록도 없음


def test_translation_cache_reuses_results(app, monkeypatch):
    """같은 문장은 두 번째부터 AI를 부르지 않음 (일괄 번역은 새 줄만 보냄)."""
    from tabs.compliance import translate
    calls = []

    def fake_call(system, content, schema, what):
        calls.append(content)
        if "items" in schema.get("properties", {}):
            lines = [json.loads(l) for l in content.split("\n")[1:]]
            return {"items": [{"i": x["i"], "text": "EN:" + x["ko"], "note_ko": ""} for x in lines]}, "claude-test"
        return {"text": "번역", "note_ko": ""}, "claude-test"

    monkeypatch.setattr(translate, "_call", fake_call)
    monkeypatch.setattr(translate, "available", lambda: True)
    with app.app_context():
        assert translate.to_korean("Say goodbye.")["text"] == "번역"
        again = translate.to_korean("Say goodbye.")
        assert again["cached"] is True and len(calls) == 1
        first = translate.to_english_batch(["가", "나"])
        assert [x["text"] for x in first["items"]] == ["EN:가", "EN:나"] and len(calls) == 2
        second = translate.to_english_batch(["가", "나", "다"])
        assert second["cached"] == 2 and len(calls) == 3 and '"다"' in calls[-1] and '"가"' not in calls[-1]
        assert translate.to_english("가")["cached"] is True and len(calls) == 3


def test_translation_survives_cache_save_failure(app, monkeypatch):
    """캐시 파일 저장이 실패해도(Windows 파일 잠금 등) 번역 결과는 정상 반환 · 메모리에는 남아 다시 부르지 않음."""
    from tabs.compliance import translate
    calls = []

    def fake_call(system, content, schema, what):
        calls.append(content)
        lines = [json.loads(l) for l in content.split("\n")[1:]]
        return {"items": [{"i": x["i"], "text": "EN:" + x["ko"], "note_ko": ""} for x in lines]}, "claude-test"

    def locked(self, target):
        raise PermissionError("[WinError 5] 액세스가 거부되었습니다")

    monkeypatch.setattr(translate, "_call", fake_call)
    monkeypatch.setattr(translate, "available", lambda: True)
    monkeypatch.setattr(translate.time, "sleep", lambda s: None)
    monkeypatch.setattr(Path, "replace", locked)
    with app.app_context():
        out = translate.to_english_batch(["라", "마"])
        assert [x["text"] for x in out["items"]] == ["EN:라", "EN:마"]
        assert translate.to_english_batch(["라"])["cached"] == 1 and len(calls) == 1


def test_only_fda_urls_are_kept_or_fetched(app):
    """업로드 파일 속 임의 주소로 서버가 내부망을 읽지 않게 (SSRF 방지)."""
    ok = "https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/a-612345-10012026"
    assert warning_letters.safe_url(ok) == ok
    assert warning_letters.safe_url("https://datadashboard.fda.gov/oii/cd/firm.htm")
    for bad in ("https://127.0.0.1/internal", "http://www.fda.gov/x", "https://www.fda.gov.evil.example/x",
                "https://user@www.fda.gov/x", "https://www.fda.gov:8443/x", "file:///etc/passwd", "https://[::1/"):
        assert warning_letters.safe_url(bad) == "", bad
    assert warning_letters.safe_url("https://datadashboard.fda.gov/x", {warning_letters.LETTER_HOST}) == ""
    wb = Workbook()
    ws = wb.active
    ws.append(["Posted Date", "Letter Issue Date", "Company Name", "Issuing Office", "Subject", "URL"])
    recent = (date.today() - timedelta(days=5)).strftime("%m/%d/%Y")
    ws.append([recent, recent, "Evil Foods", "Human Foods Program", "Misbranded", "https://127.0.0.1/internal"])
    buf = io.BytesIO()
    wb.save(buf)
    rows, _ = warning_letters.parse_xlsx(buf.getvalue())
    assert rows and rows[0]["url"] == ""
    with app.app_context():
        try:
            warning_letters._fetch("https://127.0.0.1/internal")
            raise AssertionError("내부 주소를 읽으면 안 됨")
        except warning_letters.LetterError:
            pass


def test_failed_source_keeps_previous_items(app, monkeypatch):
    """한 출처가 실패해도 지난번 정상 자료를 지우지 않고 stale로 표시 · 출처별 갱신 시각."""
    def ok():
        return [{"id": "r1", "date": "2026-10-01", "title": "Jelly recall", "summary": ""}]

    def boom():
        raise RuntimeError("일시 오류")
    with app.app_context():
        monkeypatch.setattr(refresh, "SOURCES", [("openfda-food-recalls", "리콜", ok)])
        first = refresh.collect()
        refresh.write(first)
        assert first["sources"][0]["updatedAt"]
        monkeypatch.setattr(refresh, "SOURCES", [("openfda-food-recalls", "리콜", boom)])
        second = refresh.collect()
        src = second["sources"][0]
        assert src["stale"] is True and [i["id"] for i in src["items"]] == ["r1"]
        assert src["updatedAt"] == first["sources"][0]["updatedAt"]
        assert second["errors"] and "일시 오류" in second["errors"][0]["error"]
        refresh.write(first)
        before = refresh._current()["generatedAt"]
        after = refresh.update_source({"id": "fda-warning-letters", "label": "경고장", "items": []})
        assert after["generatedAt"] == before   # 출처 하나만 바꾸면 전체 수집 시각은 그대로
        assert next(s for s in after["sources"] if s["id"] == "fda-warning-letters")["updatedAt"]


def test_ai_usage_records_cost_and_budget_stops(app, monkeypatch):
    """AI 호출마다 토큰·예상 비용 기록 · 이번 달 한도를 넘으면 AI 기능만 멈춤 (GPT 검토 ⑭)."""
    from tabs.compliance import ai_usage
    res = SimpleNamespace(model="claude-opus-5-5", usage=SimpleNamespace(
        input_tokens=1_000_000, output_tokens=100_000, cache_creation_input_tokens=0, cache_read_input_tokens=0))
    with app.app_context():
        monkeypatch.setenv("AI_MONTHLY_BUDGET_USD", "10")
        assert ai_usage.over_budget() == ""
        ai_usage.record("compliance", "AI 2차 판정", res)          # 4 + 2 = 6달러
        s = ai_usage.summary()
        assert s["calls"] == 1 and abs(s["usd"] - 6.0) < 1e-6 and s["features"][0]["feature"] == "AI 2차 판정"
        assert ai_usage.over_budget() == ""
        ai_usage.record("cs", "CS 답장 초안", res)                 # 12달러 > 10달러
        assert "한도" in ai_usage.over_budget()
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        try:
            ai_review.review("Tastes like lime.")
            raise AssertionError("한도를 넘으면 AI를 부르지 않아야 함")
        except ai_review.ReviewError as e:
            assert e.status == 429
        monkeypatch.setenv("AI_MONTHLY_BUDGET_USD", "0")         # 0 = 한도 없음
        assert ai_usage.over_budget() == ""


def test_ai_review_has_few_shot_without_eval_sentences():
    """few-shot 예시는 시스템 프롬프트에 있고, '처음 보는 문장 시험' 문장은 예시에 넣지 않음 (시험이 외운 결과가 되지 않게)."""
    import re
    src = (Path(__file__).parent.parent / "static" / "data.js").read_text(encoding="utf-8")
    block = src[src.index("evalSet: ["):src.index("],", src.index("evalSet: ["))]
    evals = re.findall(r"text: '([^']+)'", block)
    assert len(evals) >= 25
    assert ai_review.FEW_SHOT in ai_review.SYSTEM and ai_review.SYSTEM.startswith(ai_review.SYSTEM_BASE)
    assert not [t for t in evals if t in ai_review.FEW_SHOT]
    assert {"checks", "uncertain"} <= set(ai_review.SCHEMA["required"])


GRAS_SAMPLE = (
    "Downloaded from FDA GRAS Notices; Last updated 10/21/2026.\r\n\r\n"
    "GRAS Notice (GRN) No.,Substance,Intended Use,Basis,Notifier,Date of closure,FDA's Letter\r\n"
    '=T("988"), Lactobacillus fermentum LfQi6, Probiotic in foods, Scientific procedures, Acme, 03/28/2022, FDA has no questions (in PDF)\r\n'
    '=T("820"), Lactobacillus fermentum CECT5716, Infant formula, Scientific procedures, Bio, 04/03/2019,'
    " At the notifier's request, FDA ceased to evaluate this notice\r\n"
    '=T("400"), D-psicose, Sweetener, Scientific procedures, Sweet Co, 06/18/2012, FDA has no questions\r\n'
).encode("cp1252")


def test_gras_notices_parse_and_match_ingredients(app, monkeypatch):
    """FDA GRAS 통지 (2026-10-25): 내려받기 파일 파싱 · 원료별 통과/전체 건수 · 통지 없는 원료는 0건 · 한글 요약 안 함."""
    rows = refresh._parse_gras(GRAS_SAMPLE)
    assert [r["grn"] for r in rows] == ["988", "820", "400"] and rows[0]["closed"] == "2022-03-28"
    monkeypatch.setattr(refresh, "_gras_rows", lambda: (rows, "2026-10-25T00:00:00Z", ""))
    with app.app_context():
        items, warns, note, extra = refresh.fda_gras_notices()
    c = extra["counts"]
    assert c["l-fermentum"] == {"label": "L. fermentum", "total": 2, "ok": 1}
    assert c["allulose"]["ok"] == 1 and c["hovenia"]["total"] == 0 and extra["noKo"] is True
    assert [i["grn"] for i in items if i["ingredient"] == "l-fermentum"] == ["988", "820"]   # 통과한 통지 먼저
    assert all(i["url"].startswith("https://www.hfpappexternal.fda.gov/") for i in items) and "헛개 0/0" in note


def test_gras_terms_match_ingredient_table():
    """refresh.GRAS_TERMS의 원료 id는 모두 data.js 원료 표에 있어야 함 (화면 칸이 비지 않게)."""
    import re
    src = (Path(__file__).parent.parent / "static" / "data.js").read_text(encoding="utf-8")
    ids = set(re.findall(r"\{ id: '([a-z0-9-]+)', status: '(?:confirmed|candidate|avoid|base)'", src))
    assert ids and not set(refresh.GRAS_TERMS) - ids


def test_review_request_link_and_reply(app, auth_client, monkeypatch):
    """문구 검토 요청 (2026-10-25): 팀별 1회용 링크 · 로그인 없이 회신 · 상태가 목록에 보임 · 토큰은 저장 안 함 · 회신은 Slack(선택)."""
    from tabs.compliance import review_requests as rr
    sent = []
    monkeypatch.setattr(rr, "_slack", lambda text: sent.append(text) or "")
    res = auth_client.post("/api/compliance/review-requests", json={
        "text": "Reboot Jelly Original: lime, lemon and cherry jellies.", "ko": "라임·레몬·체리 젤리", "channel": "b2c",
        "product": "jelly", "aiVerdict": "PASS", "teams": ["legal", "quality", "nope"], "memo": "이번 주 인스타 게시용"})
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert set(out["links"]) == {"legal", "quality"} and out["request"]["no"].startswith("CR-")
    assert out["request"]["ruleVerdict"] == "PASS" and "tokenHash" not in json.dumps(out["request"])
    assert not sent   # 요청 알림은 Slack으로 안 감 (2026-10-26 · 엑셀 메일로)
    token = out["links"]["legal"].rsplit("/", 1)[1]
    raw = (Path(app.config["DATA_DIR"]) / "compliance" / "review-requests.json").read_text(encoding="utf-8")
    assert token not in raw   # 링크 토큰은 지문만 저장

    anon = app.test_client()   # 로그인 없이 회신
    page = anon.get(f"/review/{token}")
    assert page.status_code == 200 and "lime, lemon and cherry" in page.get_data(as_text=True)
    bad = anon.post(f"/review/{token}", data={"decision": "반려", "name": "김법무", "comment": ""})
    assert bad.status_code == 400 and "이유" in bad.get_data(as_text=True)
    ok = anon.post(f"/review/{token}", data={"decision": "승인", "name": "김법무", "comment": ""})
    assert ok.status_code == 200 and "회신을 보냈습니다" in ok.get_data(as_text=True)
    assert anon.get("/review/not-a-real-token").status_code == 404
    assert anon.get("/api/compliance/review-requests").status_code == 401   # 목록은 로그인 필요
    assert anon.get(f"/api/compliance/review-requests/{out['request']['id']}.xlsx").status_code == 401
    items = auth_client.get("/api/compliance/review-requests").get_json()["items"]
    teams = items[0]["teams"]
    assert teams["legal"]["status"] == "승인" and teams["legal"]["by"] == "김법무" and teams["quality"]["status"] == "대기"
    assert "*[REBOOT 문구 검토 회신]* 법무팀 *승인*" in sent[-1]
    assert items[0]["ruleLegal"] == ""   # PASS 문구는 법적 신호 없음
    # 웹훅이 없으면 회신 알림은 조용히 건너뜀
    monkeypatch.setattr(rr, "env", lambda k: "")
    assert rr._slack("x") == ""


def test_review_request_xlsx_lists_hits(auth_client):
    """엑셀 요청서 (2026-10-26): 요청서 · 걸린 표현 · 팀 회신 3장 · 걸린 규칙은 서버 재검사 기준 · 화면 설명은 걸린 규칙 것만 · 링크는 안 들어감."""
    from openpyxl import load_workbook
    res = auth_client.post("/api/compliance/review-requests", json={
        "text": "Clinically proven formula.", "channel": "b2c", "product": "jelly", "teams": ["legal", "regulatory"],
        "details": {"R5": {"say": "시험 결과를 말하려면 근거가 있어야 해요", "todo": ["광고에서는 빼기"], "law": ["FTC 16 CFR 255"]},
                    "R1": {"say": "지어낸 설명"}}})
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert out["request"]["ruleLegal"] == "yellow"
    x = auth_client.get(f"/api/compliance/review-requests/{out['request']['id']}.xlsx")
    assert x.status_code == 200 and "attachment" in x.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(x.data))
    assert wb.sheetnames == ["요청서", "걸린 표현", "팀 회신"]
    cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    joined = "\n".join(cells)
    assert "Clinically proven formula." in joined and "법무 확인" in joined
    assert "시험 결과를 말하려면" in joined and "FTC 16 CFR 255" in joined
    assert "지어낸 설명" not in joined   # 서버가 안 잡은 규칙(R1) 설명은 버림
    assert not any(u.rsplit("/", 1)[1] in joined for u in out["links"].values())
    assert [wb["팀 회신"].cell(row=r, column=1).value for r in (5, 6)] == ["법무팀", "인허가팀"]
    assert auth_client.get("/api/compliance/review-requests/nope.xlsx").status_code == 404


def test_review_request_mail_attaches_xlsx(auth_client, monkeypatch):
    """메일 보내기: 등록된 담당자에게만 · 엑셀 첨부 · 본문 끝에 팀별 회신 링크 · 남의 링크는 거절 · 팀을 안 고르면 요청 거절."""
    import base64
    import email
    from tabs.cs import contacts, gmail
    monkeypatch.setattr(contacts, "load", lambda: [{"role": "법무 담당", "name": "김법무", "email": "legal@example.com"}])
    sent = []

    class Svc:
        def users(self):
            return self

        def messages(self):
            return self

        def send(self, userId, body):
            sent.append(body["raw"])
            return SimpleNamespace(execute=lambda: {"id": "m1"})
    monkeypatch.setattr(gmail, "service", lambda interactive=False: Svc())
    setup = auth_client.get("/api/compliance/review-requests/mail-setup").get_json()
    assert setup["contacts"][0]["team"] == "legal"
    out = auth_client.post("/api/compliance/review-requests", json={"text": "Restart Your Day.", "teams": ["legal"]}).get_json()
    url = f"/api/compliance/review-requests/{out['request']['id']}/mail"
    body = {"to": ["legal@example.com"], "subject": "[문구 검토 요청] 테스트", "body": "검토 부탁드립니다.", "links": out["links"]}
    assert auth_client.post(url, json=body | {"to": ["x@example.com"]}).status_code == 400
    assert auth_client.post(url, json=body | {"links": {"legal": "http://x/review/fake"}}).status_code == 400
    res = auth_client.post(url, json=body)
    assert res.status_code == 200, res.get_json()
    msg = email.message_from_bytes(base64.urlsafe_b64decode(sent[0]))
    parts = list(msg.walk())
    text = next(p for p in parts if p.get_content_type() == "text/plain").get_payload(decode=True).decode("utf-8")
    assert out["links"]["legal"] in text and "검토 부탁드립니다." in text
    att = next(p for p in parts if p.get_filename())
    assert att.get_filename().endswith(".xlsx") and att.get_payload(decode=True)[:2] == b"PK"
    items = auth_client.get("/api/compliance/review-requests").get_json()["items"]
    assert items[0]["mail"]["to"] == ["legal@example.com"]
    assert auth_client.post("/api/compliance/review-requests", json={"text": "Hi", "teams": []}).status_code == 400

def test_review_request_preview_edit_and_import_replies(auth_client):
    """보내기 전 미리 받기(저장 안 함) → 고쳐서 첨부(요청번호가 들어감) → 팀이 [팀 회신]을 채워 돌려준 엑셀을 올리면 상태 반영."""
    import base64
    from openpyxl import load_workbook
    body = {"text": "Clinically proven formula.", "channel": "b2c", "product": "jelly", "teams": ["legal", "quality"]}
    pv = auth_client.post("/api/compliance/review-requests/preview", json=body)
    assert pv.status_code == 200 and pv.data[:2] == b"PK"
    assert auth_client.get("/api/compliance/review-requests").get_json()["items"] == []   # 미리보기는 저장 안 함
    wb = load_workbook(io.BytesIO(pv.data))
    assert wb["요청서"]["B4"].value == "(보내면 붙음)"
    wb["걸린 표현"]["G5"] = "마케터가 덧붙인 설명"
    buf = io.BytesIO()
    wb.save(buf)
    assert auth_client.post("/api/compliance/review-requests", json=body | {"attachment": "!!"}).status_code == 400
    assert auth_client.post("/api/compliance/review-requests",
                            json=body | {"attachment": base64.b64encode(b"not xlsx").decode()}).status_code == 400
    out = auth_client.post("/api/compliance/review-requests", json=body | {"attachment": base64.b64encode(buf.getvalue()).decode()}).get_json()
    x = auth_client.get(f"/api/compliance/review-requests/{out['request']['id']}.xlsx")
    sent = load_workbook(io.BytesIO(x.data))
    assert sent["요청서"]["B4"].value == out["request"]["no"] and sent["걸린 표현"]["G5"].value == "마케터가 덧붙인 설명"

    # 팀이 엑셀로 회신
    ws = sent["팀 회신"]
    ws["B5"], ws["D5"] = "승인", "김법무"
    ws["B6"] = "반려"   # 이유 없음 → 건너뜀
    buf = io.BytesIO()
    sent.save(buf)
    res = auth_client.post("/api/compliance/review-requests/import", data={"file": (io.BytesIO(buf.getvalue()), "reply.xlsx")},
                           content_type="multipart/form-data")
    assert res.status_code == 200, res.get_json()
    got = res.get_json()
    assert got["updated"] == ["legal"] and got["skipped"] and "의견" in got["skipped"][0]
    assert got["request"]["teams"]["legal"]["status"] == "승인" and got["request"]["teams"]["quality"]["status"] == "대기"
    bad = auth_client.post("/api/compliance/review-requests/import", data={"file": (io.BytesIO(b"hello"), "a.xlsx")},
                           content_type="multipart/form-data")
    assert bad.status_code == 400
